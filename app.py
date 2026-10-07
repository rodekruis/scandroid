from flask import (
    Flask,
    render_template,
    request,
    send_file,
    redirect,
    session,
    url_for,
    flash,
    jsonify,
    send_from_directory,
)
import requests
from io import BytesIO
from config_loader import load_config, save_config
import json
import os
from config import ADMIN_USERNAME, ADMIN_PASSWORD, FSP_USERNAME, FSP_PASSWORD
from urllib.parse import quote
import subprocess
import zipfile
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A5, landscape
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
import qrcode
import csv
from flask_session import Session
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from openpyxl import load_workbook
import sys
from config_loader import load_display_config, save_display_config

# Register a Unicode-safe font
pdfmetrics.registerFont(
    TTFont("DejaVu", os.path.join("static", "fonts", "DejaVuSans.ttf"))
)

# ---------------------------------------------------------------------------
# Arabic (right-to-left) text support for PDF rendering
#
# ReportLab's drawString does NOT perform Arabic contextual shaping (joining
# letters into their initial/medial/final forms) or bidirectional reordering.
# Without the two steps below, Arabic prints as isolated, reversed letters.
#   1. arabic_reshaper -> connects the letters into their correct joined forms
#   2. python-bidi (get_display) -> reorders logical text into visual RTL order
# Install with:  pip install arabic-reshaper python-bidi
# ---------------------------------------------------------------------------
try:
    import arabic_reshaper
    from bidi.algorithm import get_display
    _HAS_ARABIC_SHAPING = True
except ImportError:
    _HAS_ARABIC_SHAPING = False

import re as _re

# Arabic + Arabic Supplement/Extended + Presentation Forms A/B ranges.
_ARABIC_RE = _re.compile(
    r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]"
)

# DejaVuSans has weak coverage of Arabic *presentation forms* (the joined
# glyphs arabic_reshaper produces), so joined Arabic can render as boxes.
# Prefer a proper Arabic font if one is present in static/fonts; otherwise
# fall back to DejaVu (better than nothing) and log a warning once.
ARABIC_FONT = "DejaVu"
for _candidate in (
    "NotoNaskhArabic-Regular.ttf",
    "NotoSansArabic-Regular.ttf",
    "Amiri-Regular.ttf",
    "Cairo-Regular.ttf",
):
    _fp = os.path.join("static", "fonts", _candidate)
    if os.path.exists(_fp):
        try:
            pdfmetrics.registerFont(TTFont("ArabicFont", _fp))
            ARABIC_FONT = "ArabicFont"
            break
        except Exception:
            pass


def _contains_arabic(text):
    """True if the string contains any Arabic-script character."""
    return bool(text) and bool(_ARABIC_RE.search(str(text)))


def _shape_rtl(text):
    """Return text ready for ReportLab drawing.

    Arabic runs are reshaped + bidi-reordered so they print correctly;
    Latin/other text passes through unchanged. Never raises — on any error
    it returns the original text so a voucher still prints.
    """
    text = "" if text is None else str(text)
    if _HAS_ARABIC_SHAPING and _contains_arabic(text):
        try:
            return get_display(arabic_reshaper.reshape(text))
        except Exception:
            return text
    return text


def _font_for(text):
    """Pick the Arabic-capable font for Arabic text, DejaVu otherwise."""
    return ARABIC_FONT if _contains_arabic(text) else "DejaVu"


# ---------------------------------------------------------------------------
# Mixed-script (Arabic + Latin) drawing
#
# _font_for() above picks ONE font for a whole string, which is wrong for any
# string that mixes scripts: "Name: محمد" would be drawn entirely in Amiri, so
# the Latin label picked up Amiri's Latin glyphs while a pure-Latin row on the
# same voucher used DejaVu. Side by side that reads as two different typefaces.
#
# The helpers below instead:
#   1. reshape + bidi-reorder the WHOLE string once, giving visual order;
#   2. split that visual string into Arabic-script vs everything-else runs;
#   3. draw each run left-to-right in its own font.
#
# Because step 1 already produced visual order, drawing the runs sequentially
# left-to-right is correct for both LTR and RTL base directions. Latin text
# therefore always renders in LATIN_FONT, no matter what it sits next to.
# ---------------------------------------------------------------------------
LATIN_FONT = "DejaVu"

# Strong-direction detection, used to pick the bidi base direction. The first
# strong character decides: "شارع Main" is an RTL paragraph, "Main شارع" an LTR
# one, and getting this wrong misplaces trailing punctuation.
_RTL_STRONG_RE = _re.compile(r"[֐-׿؀-ۿ܀-ݏݐ-ݿࢠ-ࣿיִ-﷿ﹰ-﻿]")
_LTR_STRONG_RE = _re.compile(r"[A-Za-zÀ-ʯͰ-ϿЀ-ӿ]")


def _base_is_rtl(text):
    """True when the first strong-direction character is RTL."""
    rtl = _RTL_STRONG_RE.search(text)
    if not rtl:
        return False
    ltr = _LTR_STRONG_RE.search(text)
    return True if not ltr else rtl.start() < ltr.start()


def _visual_text(text):
    """Reshape + bidi-reorder a string into the visual order to be drawn."""
    text = "" if text is None else str(text)
    if not (_HAS_ARABIC_SHAPING and _contains_arabic(text)):
        return text
    try:
        reshaped = arabic_reshaper.reshape(text)
    except Exception:
        reshaped = text
    base = "R" if _base_is_rtl(text) else "L"
    try:
        return get_display(reshaped, base_dir=base)
    except TypeError:
        # Older/newer python-bidi builds without the base_dir keyword.
        try:
            return get_display(reshaped)
        except Exception:
            return text
    except Exception:
        return text


def _script_runs(visual):
    """Split visual-order text into [(run_text, font_name), ...].

    Neutral characters (spaces, digits, punctuation) attach to the run before
    them so a single word is never split across two fonts; leading neutrals
    attach to the run that follows.
    """
    if not visual:
        return []

    runs = []          # [[chars], is_arabic]
    pending = []       # neutrals waiting for a run to attach to

    for ch in visual:
        if _ARABIC_RE.match(ch):
            is_arabic = True
        elif _LTR_STRONG_RE.match(ch):
            is_arabic = False
        else:
            pending.append(ch)
            continue

        if runs and runs[-1][1] == is_arabic:
            runs[-1][0].extend(pending)
            runs[-1][0].append(ch)
        else:
            if runs:
                runs[-1][0].extend(pending)   # trailing neutrals stay behind
                runs.append([[ch], is_arabic])
            else:
                runs.append([pending + [ch], is_arabic])
        pending = []

    if pending:
        if runs:
            runs[-1][0].extend(pending)
        else:
            runs.append([pending, False])     # no strong chars at all

    return [
        ("".join(chars), ARABIC_FONT if is_arabic else LATIN_FONT)
        for chars, is_arabic in runs
    ]


def _mixed_width(c, text, size):
    """Drawn width of text once split into per-script runs."""
    return sum(
        c.stringWidth(run, font, size) for run, font in _script_runs(_visual_text(text))
    )


def _draw_mixed(c, x, y, text, size):
    """Draw text left-aligned at x, each script run in its own font.

    Returns the total width drawn.
    """
    cursor = x
    for run, font in _script_runs(_visual_text(text)):
        c.setFont(font, size)
        c.drawString(cursor, y, run)
        cursor += c.stringWidth(run, font, size)
    return cursor - x


def _draw_mixed_right(c, right_x, y, text, size):
    """Draw text so its right edge sits at right_x, per-script fonts intact."""
    runs = _script_runs(_visual_text(text))
    total = sum(c.stringWidth(run, font, size) for run, font in runs)
    cursor = right_x - total
    for run, font in runs:
        c.setFont(font, size)
        c.drawString(cursor, y, run)
        cursor += c.stringWidth(run, font, size)
    return total


def _draw_mixed_centred(c, cx, y, text, size):
    """Draw text centred on cx, each script run in its own font."""
    runs = _script_runs(_visual_text(text))
    total = sum(c.stringWidth(run, font, size) for run, font in runs)
    cursor = cx - total / 2.0
    for run, font in runs:
        c.setFont(font, size)
        c.drawString(cursor, y, run)
        cursor += c.stringWidth(run, font, size)
    return total


def _fit_mixed(c, text, size, max_width, min_size=7.0):
    """Largest font size <= size at which text fits max_width.

    Returns (size, text). If the text still does not fit at min_size it is
    truncated with an ellipsis so a long value can never run off the voucher.
    """
    if max_width <= 0:
        return size, text
    while size > min_size and _mixed_width(c, text, size) > max_width:
        size -= 0.5
    if _mixed_width(c, text, size) <= max_width:
        return size, text

    truncated = str(text)
    while truncated and _mixed_width(c, truncated + "…", size) > max_width:
        truncated = truncated[:-1]
    return size, (truncated + "…") if truncated else ""


app = Flask(__name__)
# The session now carries each operator's 121 access token. Flask-Session keeps
# the session data server-side (filesystem) and the cookie holds only the signed
# session id -- but that id is the bearer, so it must be signed with a real key.
# "your_secret_key" was a literal in the repo. Falling back to a random key is
# fine locally; on Azure each gunicorn worker would generate a different one and
# staff would appear to be logged out at random, so warn loudly.
app.secret_key = os.getenv("FLASK_SECRET_KEY") or os.urandom(32).hex()
if not os.getenv("FLASK_SECRET_KEY"):
    print(
        "[auth] WARNING: FLASK_SECRET_KEY is not set - a random key was generated. "
        "Set it in Azure App Settings or logins will drop unpredictably."
    )
app.config["SESSION_TYPE"] = "filesystem"
app.config["SESSION_PERMANENT"] = False
Session(app)


@app.context_processor
def inject_national_society():
    config = load_config()
    return {"national_society": config.get("nationalSociety", "")}


def _make_qr_image(data, box_cm=3.0):
    """Return a Pillow image for the QR sized to box_cm × box_cm at 300dpi."""
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=4,
    )
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white").convert("RGB")

    # resize to cm at 300dpi (ReportLab draws images in points, we’ll scale when drawing)
    target_px = int((box_cm / 2.54) * 300)  # cm -> inches -> px
    img = img.resize((target_px, target_px))
    return img


# ---------------------------------------------------------------------------
# Instance-static helpers (used by the voucher designer + voucher rendering)
#
# Per-program voucher logos are stored in the same instance static folder that
# /instance-static/<filename> serves from. On Azure that is the per-context
# config folder; locally it falls back to the app's static/ folder.
# ---------------------------------------------------------------------------
def _instance_static_dir():
    env = os.getenv("SCANDROID_ENV", "local")
    context = os.getenv("SCANDROID_CONTEXT", "local")
    if env == "azure":
        path = f"/home/site/configs/{context}/static"
    else:
        path = os.path.join(app.root_path, "static")
    os.makedirs(path, exist_ok=True)
    return path


def _resolve_logo_path(filename, static_folder=None):
    """Return the on-disk path for a stored logo filename, or None."""
    if not filename:
        return None
    bases = [_instance_static_dir()]
    if static_folder:
        bases.append(static_folder)
    for base in bases:
        candidate = os.path.join(base, filename)
        if os.path.exists(candidate):
            return candidate
    return None


def _default_label(raw):
    """Default label for a column: the header exactly as the file spells it.

    Only surrounding/repeated whitespace is collapsed — underscores, casing and
    script are left alone, so "bank_account_number" prints as
    "bank_account_number". Admins who want something tidier type their own label
    on the Generate Vouchers page, which is stored per program and never
    overwritten by this.
    """
    return " ".join(str(raw or "").split())


def _is_reference_id(key):
    """True for any spelling of the reference-ID column (it is not printed as
    a field — it becomes the QR code and the small line at the page foot)."""
    k = str(key or "").strip().lower().replace("_", "").replace(" ", "")
    return "ref" in k and "id" in k


def _voucher_field_config(design):
    """The saved [{key, label, show, custom}] field list for a program, or []."""
    fields = (design or {}).get("fields")
    return [f for f in fields if isinstance(f, dict)] if isinstance(fields, list) else []


# An admin-typed label is capped so a pasted paragraph cannot wreck the layout.
MAX_FIELD_LABEL_LEN = 60


def _clean_field_label(value, key, default=None):
    """Sanitise a label, falling back to the column's default label."""
    label = " ".join(str(value or "").split())[:MAX_FIELD_LABEL_LEN].strip()
    return label or default or _default_label(key)


def merge_voucher_fields(design, columns):
    """Reconcile the columns found in an upload with the program's saved order.

    columns: [(key, header)] in the order they appear in the file.

    Saved fields that are still present keep their saved position and
    visibility; columns the program has never seen are appended in file order
    and default to visible.

    Labels: a field the admin has renamed (custom=True) keeps that label
    forever. Everything else shows the column header verbatim, so fixing a
    header in the CSV still flows through to the voucher.

    Each entry also carries "default_label" — the label this column WOULD have
    with no override — so the UI can offer a reset. It is derived, not stored.
    """
    default_by_key, order = {}, []
    for key, header in columns:
        key = str(key or "").strip().lower()
        if not key or _is_reference_id(key) or key in default_by_key:
            continue
        # The header exactly as written; fall back to the key if the header
        # cell was blank.
        default_by_key[key] = _default_label(header) or _default_label(key)
        order.append(key)

    merged, used = [], set()
    for entry in _voucher_field_config(design):
        key = str(entry.get("key") or "").strip().lower()
        if key not in default_by_key or key in used:
            continue
        used.add(key)
        default_label = default_by_key[key]
        is_custom = bool(entry.get("custom"))
        merged.append(
            {
                "key": key,
                "label": _clean_field_label(entry.get("label"), key, default_label)
                if is_custom
                else default_label,
                "custom": is_custom,
                "default_label": default_label,
                "show": entry.get("show") is not False,
            }
        )

    for key in order:
        if key not in used:
            merged.append(
                {
                    "key": key,
                    "label": default_by_key[key],
                    "custom": False,
                    "default_label": default_by_key[key],
                    "show": True,
                }
            )

    return merged


def _ordered_fields(item, design):
    """[(label, value)] to print on one voucher, in the configured order.

    Hidden fields, the reference ID and empty values are filtered out. Columns
    with no saved configuration print last, in file order, so a new CSV column
    never silently disappears from the voucher.
    """
    printable = {
        str(k).strip().lower(): v
        for k, v in (item or {}).items()
        if k and not _is_reference_id(k) and v not in (None, "")
    }

    out, used = [], set()
    for entry in _voucher_field_config(design):
        key = str(entry.get("key") or "").strip().lower()
        if key not in printable or key in used:
            continue
        used.add(key)
        if entry.get("show") is False:
            continue
        out.append((entry.get("label") or _default_label(key), printable[key]))

    for key, value in printable.items():
        if key not in used:
            out.append((_default_label(key), value))

    return out


def _draw_voucher(c, item, static_folder, design=None):
    """
    Draw one voucher on an A5 LANDSCAPE page matching the provided layout.
    Supports dynamic CSV fields.
    """
    width, height = landscape(A5)

    margin = 1.0 * cm
    inner_w = width - 2 * margin
    inner_h = height - 2 * margin

    # ---- BORDER -------------------------------------------------------------
    c.setLineWidth(0.8)
    c.rect(margin, margin, inner_w, inner_h)
    logo_top_y = height - margin - 1.0 * cm  # move logos DOWN slightly (was -0.0)

    # Per-program design (title / subtitle / logos / sizes), with hardcoded fallbacks
    design = design or {}
    logo_height = (
        LEFT_LOGO_HEIGHT_CM.get(_clean_size(design.get("logo1_size")), 2.5) * cm
    )
    left_logo_path = _resolve_logo_path(
        design.get("logo1"), static_folder
    ) or os.path.join(static_folder, "ns1.png")
    right_logo_path = _resolve_logo_path(
        design.get("logo2"), static_folder
    ) or os.path.join(static_folder, "ns2.png")

    # LEFT LOGO – scale by HEIGHT, derive width from aspect ratio.
    # We must pass an explicit width: with preserveAspectRatio=True and only a
    # height, ReportLab has no bounding box and wide images fail to render.
    left_logo_top = height - margin - 0.5*cm
    try:
        logo1 = ImageReader(left_logo_path)
        img_w, img_h = logo1.getSize()
        scale = logo_height / img_h
        scaled_width = img_w * scale
        c.drawImage(
            logo1,
            margin,                          # as far left as possible
            logo_top_y - logo_height,
            width=scaled_width,
            height=logo_height,
            preserveAspectRatio=True,
            mask='auto'
        )
        left_logo_top = logo_top_y
    except Exception as e:
        app.logger.warning("Left voucher logo failed to draw (%s): %s", left_logo_path, e)

    # RIGHT LOGO – wide logo, scale by width
    right_logo_top = height - margin - 0.5*cm
    try:
        logo2 = ImageReader(right_logo_path)
        max_width = (
            RIGHT_LOGO_WIDTH_CM.get(_clean_size(design.get("logo2_size")), 5.0) * cm
        )

        img_w, img_h = logo2.getSize()
        scale = max_width / img_w
        scaled_height = img_h * scale

        c.drawImage(
            logo2,
            width - margin - max_width,
            logo_top_y - scaled_height,
            width=max_width,
            height=scaled_height,
            preserveAspectRatio=True,
            mask="auto",
        )
        right_logo_top = logo_top_y
    except Exception as e:
        app.logger.warning("Right voucher logo failed to draw (%s): %s", right_logo_path, e)

    # --- CENTER "Project" between logos (vertically aligned to logos) ---
    project_y = min(left_logo_top, right_logo_top) - 0.3 * cm

    c.setFont("DejaVu", 16)
    c.drawCentredString(width / 2, project_y, "Project")

    # underline "Project"
    text_width = c.stringWidth("Project", "DejaVu", 16)
    c.line(
        (width / 2 - text_width / 2),
        project_y - 0.08 * cm,
        (width / 2 + text_width / 2),
        project_y - 0.08 * cm,
    )

    # --- MAIN TITLE just below "Project" ---
    title_text = design.get("title") or "CASH ON THE MOVE"
    title_y = project_y - 1.4*cm
    # Per-run fonts: a title like "CASH ON THE MOVE / نقد" keeps its Latin half
    # in DejaVu instead of switching the whole line to the Arabic face.
    title_size, title_text = _fit_mixed(c, title_text, 22, inner_w - 1.0 * cm, min_size=11)
    _draw_mixed_centred(c, width / 2, title_y, title_text, title_size)

    # Subtitle lines (newline-separated; falls back to the original copy)
    subtitle_raw = design.get("subtitle")
    if subtitle_raw:
        subtitle_lines = [
            ln.strip() for ln in str(subtitle_raw).splitlines() if ln.strip()
        ]
    else:
        subtitle_lines = [
            "Supporting people on the move especially those",
            "in vulnerable situations",
        ]

    sub_y = title_y - 1.0*cm
    for line in subtitle_lines:
        line_size, line_text = _fit_mixed(c, line, 11, inner_w - 1.0 * cm, min_size=7)
        _draw_mixed_centred(c, width / 2, sub_y, line_text, line_size)
        sub_y -= 0.6*cm
    # ---- QR CODE ------------------------------------------------------------
    show_qr = design.get("show_qr", True)
    refid = item.get("referenceid", "").strip()
    qr_box_size = 5.0 * cm
    qr_x = margin + 0.6 * cm
    qr_y = margin + 0.6 * cm

    if show_qr:
        # Outer box
        c.setLineWidth(0.7)
        c.rect(qr_x - 0.1*cm, qr_y - 0.1*cm, qr_box_size + 0.2*cm, qr_box_size + 0.2*cm)

        # QR image
        qr_img = _make_qr_image(refid, box_cm=3.0)
        c.drawInlineImage(qr_img, qr_x, qr_y, width=qr_box_size, height=qr_box_size)

    # *** Removed “Reference ID below QR” (as you requested) ***

    # ---- BENEFICIARY INFORMATION -------------------------------------------
    if show_qr:
        info_x = qr_x + qr_box_size + 2.0*cm
    else:
        # No QR: start the info block at the left margin instead of leaving a gap.
        info_x = qr_x
    info_y = qr_y + qr_box_size - 0.5*cm
    info_right = width - margin - 0.4 * cm
    avail_w = info_right - info_x

    # Fields in the order configured for this program (see _ordered_fields):
    # [(label, value), ...]. Hidden fields and the reference ID are already
    # filtered out.
    rows_to_print = _ordered_fields(item, design)

    # ---- Vertical fit -------------------------------------------------------
    # Room between the first row's baseline and the small reference ID line.
    base_size = 12.0
    line_height = 0.75 * cm
    avail_h = info_y - (margin + 0.9 * cm)
    if rows_to_print and len(rows_to_print) * line_height > avail_h:
        line_height = avail_h / len(rows_to_print)
        # Keep the type proportional to the tightened leading, with a floor so
        # a 15-column CSV stays legible rather than collapsing to nothing.
        base_size = max(7.0, min(base_size, line_height / cm * 16.0))

    # ---- Label column width -------------------------------------------------
    # Every value starts at the same x, so Arabic and Latin rows line up in a
    # column instead of Arabic rows jumping to the right margin.
    label_gap = 0.35 * cm
    widest_label = 0.0
    for label, _value in rows_to_print:
        widest_label = max(widest_label, _mixed_width(c, f"{label}:", base_size))
    label_col = min(widest_label + label_gap, avail_w * 0.55)

    # ---- Block direction ----------------------------------------------------
    # The COLUMN NAMES decide which way the rows read. Latin headers give the
    # familiar left-to-right block (labels left, values in a column to their
    # right) no matter what script the values are in — which is the whole point
    # of this layout. Arabic headers mirror the block so it reads right-to-left
    # with the label first. It is decided once for the whole block, on a
    # majority of the labels, so every row aligns to the same edge even when
    # the CSV mixes Latin and Arabic headers.
    arabic_labels = sum(1 for label, _v in rows_to_print if _contains_arabic(label))
    block_rtl = bool(rows_to_print) and arabic_labels * 2 > len(rows_to_print)

    label_w = label_col - 0.1 * cm
    value_w = avail_w - label_col

    y = info_y
    for label, value in rows_to_print:
        label_size, label_text = _fit_mixed(c, f"{label}:", base_size, label_w, min_size=6)
        value_size, value_text = _fit_mixed(c, str(value), base_size, value_w, min_size=6)

        if block_rtl:
            # Label hugs the right edge, values right-aligned in the column to
            # its left, so the eye meets the label first.
            _draw_mixed_right(c, info_right, y, label_text, label_size)
            _draw_mixed_right(c, info_right - label_col, y, value_text, value_size)
        else:
            _draw_mixed(c, info_x, y, label_text, label_size)
            _draw_mixed(c, info_x + label_col, y, value_text, value_size)

        y -= line_height

    # ---- SMALL REFERENCE ID AT BOTTOM ------------------------------------
    refid = item.get("referenceid", "").strip()
    if refid:
        c.setFont("DejaVu", 8)
        c.drawString(margin + 0.1 * cm, margin + 0.2 * cm, f"Reference ID: {refid}")


def generate_vouchers_pdf(rows, static_folder, design=None):
    """
    rows: list of dicts with keys: referenceId, name
    design: optional per-program voucher design (title/subtitle/logos)
    returns BytesIO of PDF
    """
    from io import BytesIO

    pdf_io = BytesIO()

    # Create a landscape A5 page
    c = canvas.Canvas(pdf_io, pagesize=landscape(A5))

    for r in rows:
        _draw_voucher(c, r, static_folder, design=design)
        c.showPage()

    c.save()
    pdf_io.seek(0)
    return pdf_io


# ---------------------------------------------------------------------------
# Program list + voucher design storage helpers
# ---------------------------------------------------------------------------
def resolve_programs(system_config=None):
    """Return (programs, lookup).

    programs: list of {"id": str, "title": str} for the program dropdowns.
    lookup:   {program_id: raw program object}.

    Tries to resolve human-readable titles from 121 (same approach as the
    field-config page); falls back to the program id if that fails.
    """
    if system_config is None:
        try:
            system_config = load_config()
        except Exception:
            system_config = {}

    programs_raw = system_config.get("PROGRAMS", [])
    programs, lookup = [], {}
    url121 = system_config.get("url121")

    if url121 and programs_raw:
        try:
            login_resp = requests.post(
                f"{url121}/api/users/login",
                json={
                    "username": system_config.get("username121", ""),
                    "password": system_config.get("password121", ""),
                },
                timeout=10,
            )
            if login_resp.status_code == 201:
                token = login_resp.json().get("access_token_general")
                cookies = {"access_token_general": token}
                for p in programs_raw:
                    pid = p.get("programId")
                    title = str(pid)
                    try:
                        r = requests.get(
                            f"{url121}/api/programs/{pid}",
                            cookies=cookies,
                            timeout=10,
                        )
                        if r.status_code == 200:
                            titles = r.json().get("titlePortal", {}) or {}
                            title = titles.get("en") or next(
                                iter(titles.values()), title
                            )
                    except Exception:
                        pass
                    programs.append({"id": str(pid), "title": title})
                    lookup[str(pid)] = p
        except Exception as e:
            print("Program title lookup failed:", e)

    # Fallback: ids only
    if not programs:
        for p in programs_raw:
            pid = str(p.get("programId"))
            programs.append({"id": pid, "title": pid})
            lookup[pid] = p

    return programs, lookup


# Note: SVG is intentionally excluded — ReportLab's ImageReader cannot render
# SVG into the voucher PDF, so only raster formats are accepted.
ALLOWED_LOGO_EXT = {"png", "jpg", "jpeg"}

# Per-logo prominence. Values stored in the design; mapped to real dimensions
# at render time (left logo is sized by height, right logo by width).
LOGO_SIZES = {"small", "medium", "large"}
LEFT_LOGO_HEIGHT_CM = {"small": 1.8, "medium": 2.5, "large": 3.3}
RIGHT_LOGO_WIDTH_CM = {"small": 4.0, "medium": 5.0, "large": 6.2}


def _clean_size(value):
    value = (value or "").strip().lower()
    return value if value in LOGO_SIZES else "medium"


def load_voucher_designs():
    """All per-program voucher designs stored in display_config.json."""
    try:
        cfg = load_display_config()
    except Exception:
        cfg = {}
    return (cfg or {}).get("voucher_designs", {}) or {}


def get_voucher_design(program_id):
    if program_id is None:
        return None
    return load_voucher_designs().get(str(program_id))


def save_voucher_design(program_id, design):
    cfg = load_display_config() or {}
    cfg.setdefault("voucher_designs", {})
    cfg["voucher_designs"][str(program_id)] = design
    save_display_config(cfg)


def _save_logo_file(file_storage, program_id, slot):
    """Persist an uploaded logo to the instance static folder; return filename.

    Raises ValueError if the file isn't an accepted raster format.
    """
    original = (file_storage.filename or "").lower()
    ext = original.rsplit(".", 1)[-1] if "." in original else ""
    if ext not in ALLOWED_LOGO_EXT:
        raise ValueError("Unsupported logo format. Please use a PNG or JPG image.")
    filename = f"voucher_logo_{program_id}_{slot}.{ext}"
    file_storage.save(os.path.join(_instance_static_dir(), filename))
    return filename


def _logo_url(filename):
    """Served URL for a stored logo, cache-busted by file mtime.

    Logo filenames are stable per program/slot and get overwritten on re-upload,
    so without a version param browsers would keep showing the old image.
    """
    if not filename:
        return None
    url = f"/instance-static/{filename}"
    path = _resolve_logo_path(filename)
    if path:
        try:
            url += f"?v={int(os.path.getmtime(path))}"
        except OSError:
            pass
    return url


def _design_for_client(design):
    """Shape a stored design for the browser (logo filenames -> served URLs)."""
    design = design or {}
    return {
        "title": design.get("title", ""),
        "subtitle": design.get("subtitle", ""),
        "logo1_url": _logo_url(design.get("logo1")),
        "logo2_url": _logo_url(design.get("logo2")),
        "logo1_size": _clean_size(design.get("logo1_size")),
        "logo2_size": _clean_size(design.get("logo2_size")),
        "show_qr": design.get("show_qr", True),
        "fields": _voucher_field_config(design),
    }



translations = {
"en": {
    # ---- Shared chrome (header / footer / buttons on every page) ----
    "footer_dev": "Developed by 510 @ The Netherlands Red Cross",
    "footer_support": "If you need support contact jharrison@redcross.nl",
    "logout": "Logout",
    "back_to_dashboard": "Back",
    "back_to_programs": "Back to Programs",
    "go_back": "Go Back",
    "save": "Save",
    "saved_successfully": "Saved successfully",
    "failed_to_save": "Failed to save",
    "remove": "Remove",
    "submit": "Submit",
    "program": "Program",
    "program_id": "Select Program",
    "select_option": "Select...",

    # ---- home.html ----
    "home_title": "Welcome",
    "home_question": "Who are you?",
    "home_admin": "Administrator",
    "home_fsp": "Distribution Staff",

    # ---- admin_login.html / fsp_login.html (errors raised in app.py) ----
    "login": "Login for Administrator",
    "fsp_login": "Login for Distribution Staff",
    "enter_password": "Enter Password",
    "login_error": "Incorrect username or password.",
    "wrong_credentials": "Incorrect username or password.",
    "api_unreachable": "Unable to reach the login server. Please try again.",

    # ---- fsp_programs.html ----
    "program_select_title": "Select a Program",
    "program_select_subtitle": "Choose the program you want to work with",

    # ---- admin_dashboard.html ----
    "admin_dashboard_title": "Admin Dashboard",
    "admin_dashboard_subtitle": "Manage your 121 Scan configuration and voucher tools.",
    "config_system": "System Configuration",
    "config_system_desc": "Edit API keys, endpoints, and system parameters.",
    "config_display": "Configure Fields to Display",
    "config_display_desc": "Configure which fields and details are shown to FSPs.",
    "generate_vouchers": "Generate Vouchers",
    "generate_vouchers_desc": "Generate QR vouchers for printing and distribution.",
    "voucher_design": "Voucher Design",
    "voucher_design_desc": "Customise the title, subtitle and logos printed on each program's vouchers.",

    # ---- system_config.html ----
    "kobo_api_token": "API token",
    "kobo_token_hint": "Found in KoboToolbox under Account settings → Security → API token.",
    "verify_token": "Verify token",
    "verifying": "Verifying…",
    "verify_unreachable": "Could not reach the server.",
    "enter_token_first": "Enter a new token above, then verify.",
    "token_invalid": "Token not valid for this server.",
    "token_linked_to": "Token linked to",
    "program_map_hint": "Each 121 program is matched to the KoboToolbox form that collects its data.",

    # ---- config.html (field display configuration) ----
    "config_title": "Configure Fields to Display",
    "config_subtitle": "Choose which fields appear during a scan and set their labels in each language. Photo capture pulls its options from the program's KoboToolbox form.",
    "config_program_hint": "Settings below apply to the selected program",
    "config_select_or_delete": "Please select a field or delete the highlighted row(s) before saving.",
    "fields_section_title": "Displayed fields",
    "fields_section_hint": "Pick a field, then adjust its label per language. Labels prefill automatically and can be edited.",
    "field_key": "Kobo Photo Field",
    "field_name_121": "121 Field Name",
    "label_en": "Label (EN)",
    "label_fr": "Label (FR)",
    "label_ar": "Label (AR)",
    "add_field": "Add Field",
    "select_field_121": "Select field",
    "editable": "Editable",
    "editable_hint": "Turn on to let distribution staff change this field during a scan. Edits are pushed back to 121 with the payment.",
    "editable_on": "On",
    "editable_off": "Off",
    "scan_config_title": "Scan Method",
    "scan_section_hint": "Choose how beneficiaries are identified.",
    "enable_qr_scan": "Enable QR code scanning",
    "enable_qr_scan_hint": "When off, this program uses manual code entry only (no QR camera).",
    "photo_config_title": "Photo Field Configuration",
    "photo_section_hint": "Maps to an image question in the program's KoboToolbox form. Labels prefill from Kobo.",
    "enable_photo_field": "Enable photo field display",
    "enable_photo_field_hint": "When on, staff can attach a photo to the selected Kobo image field.",
    "kobo_prefill_note": "Labels filled from KoboToolbox — edit any if needed",

    # ---- scan.html ----
    "scan_title": "Scan QR",
    "scan_hint": "Point the camera at the QR code.",
    "code_entry_title": "Enter code",
    "code_entry_hint": "Enter the beneficiary's unique code below.",
    "no_qr_btn": "Can't scan QR?",
    "manual_code_label": "Enter code manually",
    "manual_code_placeholder": "Type or paste the code…",
    "manual_submit": "Go",
    "start_camera": "Start camera",
    "waiting_to_start": "Waiting to start…",
    "requesting_camera": "Requesting camera… If prompted, tap Allow.",
    "camera_denied": "Camera permission denied or not available.",
    "scanning": "Scanning…",
    "starting_camera": "Starting camera…",
    "scan_trouble_hint": "Still trying to read the code — hold steady and move a little closer.",
    "decoder_unavailable": "The scanner could not start on this device. Enter the Reference ID printed below the QR code using the box below.",
    "no_read_banner": "Can’t read this QR code",
    "no_read_banner_sub": "Type the Reference ID printed under the QR code into the box below.",
    "no_read_hint": "This QR code can’t be read — it may be blurred, faded or low quality. Enter the Reference ID printed below the QR code using the box below.",

    # ---- fsp_admin.html (connection state, counters, send) ----
    "current_program": "Program",
    "status_online": "Online",
    "status_online_sub": "You have internet",
    "status_offline": "No connection",
    "status_offline_sub": "Check your signal",
    "status_offline_mode": "Offline mode",
    "status_offline_mode_sub": "The app will not use the internet",
    "work_offline": "Work offline",
    "go_back_online": "Go back online",
    "beneficiaries_ready_to_scan": "Ready to scan:",
    "last_synced": "Last synced:",
    "cannot_sync_offline": "Cannot sync while offline",
    "people_scanned": "People scanned:",
    "total_amount": "Total amount:",
    "payments_ready_send": "Payments ready to send",
    "last_payments_submit": "Last submission:",
    "cannot_send_offline": "Cannot send while offline",
    "payment_submit_success": "✅ Payments submitted successfully!",
    "payment_submit_failed": "❌ Failed to submit",
    "no_payments": "No payments ready to send.",

    # ---- beneficiary_offline.html ----
    "title": "Beneficiary Information",
    "decryption_enabled": "Decryption enabled",
    "loading_photo": "Loading photo…",
    "payment_amount": "Payment Amount",
    "multiple_payments": "This beneficiary has more than one payment on record. Select the one you are distributing.",
    "select_payment": "Select a payment",
    "selected": "Selected",
    "all_collected": "All payments for this beneficiary have already been collected.",
    "payment_approved": "Approve",
    "payment_rejected": "Skip",
    "enter_value": "Enter value",

    # ---- success_offline.html ----
    "success_title": "Successfully submitted",
    "success_message": "You may now scan the next beneficiary.",
    "payments_ready": "Payments ready to push to 121:",
    "scan_next": "Scan next beneficiary",
    "go_home": "Finished scanning? Go to homepage",

    # ---- invalid-qr.html ----
    "invalid_qr_title": "Invalid QR Code",
    "invalid_qr_message": "This QR code cannot be used.",
    "checking_reason": "Checking reason…",
    "reason_used": "This QR code has already been used.",
    "reason_invalid": "The QR code is invalid or unrecognized.",
    "reason_no_record": "This QR code does not match any stored beneficiary.",
    "reason_database": "There was a problem reading offline data. Please try again.",

    # ---- vouchers.html (upload + field order) ----
    "voucher_generator": "Voucher Generator",
    "csv_hint": "Upload a CSV or Excel file with referenceId and any extra fields to print",
    "choose_csv": "Choose CSV or Excel file…",
    "upload_csv": "Upload CSV or Excel file",
    "download_vouchers": "Download vouchers (PDF)",
    "choose_csv_alert": "Please choose a CSV first.",
    "upload_failed": "Upload failed",
    "voucher_ready_singular": "voucher ready to download",
    "voucher_ready_plural": "vouchers ready to download",
    "voucher_design_incomplete_title": "Voucher design not completed",
    "voucher_design_incomplete_desc": "This program doesn't have a voucher design yet. Set the title, subtitle and logos before generating vouchers.",
    "voucher_design_go": "Design this voucher",
    "voucher_fields_title": "Fields on the voucher",
    "voucher_fields_hint": "Drag to reorder. Click a name to change how it is printed. Switch a field off to leave it off the voucher. This is saved for this program.",
    "voucher_fields_none": "No printable columns were found in this file.",
    "voucher_fields_order_saved": "Field order saved",
    "voucher_fields_order_failed": "Could not save field order",
    "voucher_fields_move_up": "Move up",
    "voucher_fields_move_down": "Move down",
    "voucher_fields_sample": "e.g.",
    "voucher_fields_rename": "Edit the name printed on the voucher",
    "voucher_fields_column": "Column in your file",
    "voucher_fields_reset": "reset",
    "voucher_fields_reset_title": "Reset to",

    # ---- voucher_design.html ----
    "voucher_design_title": "Voucher Design",
    "voucher_design_settings": "Design settings",
    "voucher_design_settings_desc": "These replace the values printed on the voucher. Each program keeps its own design.",
    "voucher_title_label": "Title",
    "voucher_subtitle_label": "Subtitle",
    "voucher_logos_label": "Logos",
    "voucher_left_logo": "Left logo",
    "voucher_right_logo": "Right logo",
    "voucher_logo_upload": "Upload",
    "voucher_logo_remove": "remove",
    "voucher_no_logo": "No logo",
    "voucher_logo_size": "Size",
    "voucher_size_small": "Small",
    "voucher_size_medium": "Medium",
    "voucher_size_large": "Large",
    "voucher_show_qr": "Show QR code",
    "voucher_show_qr_hint": "Turn off if this program's vouchers don't need a QR code.",
    "voucher_live_preview": "Live preview",
    "voucher_preview_desc": "A5 landscape, matching the printed voucher layout.",
    "voucher_save_design": "Save design",
    "voucher_design_saved": "Design saved",
    "voucher_design_save_failed": "Could not save design",
    "lookup_min": "Type at least {n} characters of the code.",
    "lookup_many": "{n} possible matches — keep typing.",
    "lookup_one": "One match found — press Go.",
    "lookup_none": "No match in the saved records. Check the code.",
    # ---- lookup field (config.html) ----
    "lookup_section_title": "Beneficiary lookup",
    "lookup_section_hint": "Which identifier staff type when the QR code can't be used.",
    "lookup_field_label": "Look up by",
    "lookup_field_default": "Reference ID (default)",
    "lookup_checking": "Checking whether this field is unique in 121…",
    "lookup_unique_ok": "Unique across all registrations in this program.",
    "lookup_not_unique": "Not unique: {values} value(s) are shared by {rows} registrations. Staff will not be able to identify those people by this field.",
    "lookup_has_blanks": "{n} registration(s) have no value for this field and cannot be found by it.",
    "lookup_check_failed": "Could not check uniqueness against 121 right now.",
    "lookup_partial_check": "Checked {checked} of {total} registrations — the result may be incomplete.",

    # ---- lookup field (scan.html) ----
    "manual_code_label_field": "Enter {field}",
    "lookup_duplicate": "{n} people share this value — use the Reference ID instead.",
    "verify_only_note": "Verification only — your account cannot record payments."    
},

"fr": {
    # ---- Shared chrome (header / footer / buttons on every page) ----
    "footer_dev": "Développé par 510 @ La Croix-Rouge néerlandaise",
    "footer_support": "Pour toute assistance, contactez jharrison@redcross.nl",
    "logout": "Déconnexion",
    "back_to_dashboard": "Retour",
    "back_to_programs": "Retour aux programmes",
    "go_back": "Retour",
    "save": "Enregistrer",
    "saved_successfully": "Enregistré avec succès",
    "failed_to_save": "Échec de l'enregistrement",
    "remove": "Supprimer",
    "submit": "Soumettre",
    "program": "Programme",
    "program_id": "Sélectionner le programme",
    "select_option": "Sélectionner...",

    # ---- home.html ----
    "home_title": "Bienvenue",
    "home_question": "Qui êtes-vous ?",
    "home_admin": "Administrateur",
    "home_fsp": "Personnel de distribution",

    # ---- admin_login.html / fsp_login.html (errors raised in app.py) ----
    "login": "Connexion administrateur",
    "fsp_login": "Connexion personnel de distribution",
    "enter_password": "Entrer le mot de passe",
    "login_error": "Nom d’utilisateur ou mot de passe incorrect.",
    "wrong_credentials": "Nom d’utilisateur ou mot de passe incorrect.",
    "api_unreachable": "Impossible de joindre le serveur. Veuillez réessayer.",

    # ---- fsp_programs.html ----
    "program_select_title": "Sélectionner un programme",
    "program_select_subtitle": "Choisissez le programme avec lequel vous souhaitez travailler",

    # ---- admin_dashboard.html ----
    "admin_dashboard_title": "Tableau de bord Admin",
    "admin_dashboard_subtitle": "Gérez la configuration de 121 Scan et les outils de bons.",
    "config_system": "Configuration du système",
    "config_system_desc": "Modifier les clés API, les points d'accès et les paramètres du système.",
    "config_display": "Configurer les champs à afficher",
    "config_display_desc": "Configurer les champs et informations affichés aux FSP.",
    "generate_vouchers": "Générer les bons",
    "generate_vouchers_desc": "Générer des bons QR pour impression et distribution.",
    "voucher_design": "Conception du bon",
    "voucher_design_desc": "Personnalisez le titre, le sous-titre et les logos imprimés sur les bons de chaque programme.",

    # ---- system_config.html ----
    "kobo_api_token": "Jeton API",
    "kobo_token_hint": "Se trouve dans KoboToolbox sous Paramètres du compte → Sécurité → Jeton API.",
    "verify_token": "Vérifier le jeton",
    "verifying": "Vérification…",
    "verify_unreachable": "Impossible de joindre le serveur.",
    "enter_token_first": "Saisissez un nouveau jeton ci-dessus, puis vérifiez.",
    "token_invalid": "Jeton non valide pour ce serveur.",
    "token_linked_to": "Jeton lié à",
    "program_map_hint": "Chaque programme 121 est associé au formulaire KoboToolbox qui collecte ses données.",

    # ---- config.html (field display configuration) ----
    "config_title": "Configurer les champs à afficher",
    "config_subtitle": "Choisissez les champs qui apparaissent lors d'un scan et définissez leurs libellés dans chaque langue. La capture de photo tire ses options du formulaire KoboToolbox du programme.",
    "config_program_hint": "Les paramètres ci-dessous s'appliquent au programme sélectionné",
    "config_select_or_delete": "Veuillez sélectionner un champ ou supprimer la ou les lignes en surbrillance avant d'enregistrer.",
    "fields_section_title": "Champs affichés",
    "fields_section_hint": "Choisissez un champ, puis ajustez son libellé par langue. Les libellés se remplissent automatiquement et peuvent être modifiés.",
    "field_key": "Champ photo Kobo",
    "field_name_121": "Nom du champ 121",
    "label_en": "Libellé (EN)",
    "label_fr": "Libellé (FR)",
    "label_ar": "Libellé (AR)",
    "add_field": "Ajouter un champ",
    "select_field_121": "Sélectionner un champ",
    "editable": "Modifiable",
    "editable_hint": "Activez pour permettre au personnel de distribution de modifier ce champ pendant un scan. Les modifications sont renvoyées à 121 avec le paiement.",
    "editable_on": "Activé",
    "editable_off": "Désactivé",
    "scan_config_title": "Méthode de scan",
    "scan_section_hint": "Choisissez comment les bénéficiaires sont identifiés.",
    "enable_qr_scan": "Activer le scan de code QR",
    "enable_qr_scan_hint": "Désactivé, ce programme utilise uniquement la saisie manuelle du code (pas de caméra QR).",
    "photo_config_title": "Configuration du champ photo",
    "photo_section_hint": "Correspond à une question image dans le formulaire KoboToolbox du programme. Les libellés se remplissent depuis Kobo.",
    "enable_photo_field": "Activer l'affichage du champ photo",
    "enable_photo_field_hint": "Lorsqu'activé, le personnel peut joindre une photo au champ image Kobo sélectionné.",
    "kobo_prefill_note": "Libellés remplis depuis KoboToolbox — modifiez-les si nécessaire",

    # ---- scan.html ----
    "scan_title": "Scanner un QR",
    "scan_hint": "Pointez la caméra vers le code QR.",
    "code_entry_title": "Saisir le code",
    "code_entry_hint": "Saisissez ci-dessous le code unique du bénéficiaire.",
    "no_qr_btn": "Impossible de scanner le QR ?",
    "manual_code_label": "Saisir le code manuellement",
    "manual_code_placeholder": "Tapez ou collez le code…",
    "manual_submit": "Valider",
    "start_camera": "Démarrer la caméra",
    "waiting_to_start": "En attente de démarrage…",
    "requesting_camera": "Demande d’accès à la caméra… Si demandé, touchez Autoriser.",
    "camera_denied": "Accès à la caméra refusé ou non disponible.",
    "scanning": "Analyse…",
    "starting_camera": "Démarrage de la caméra…",
    "scan_trouble_hint": "Lecture du code en cours — restez stable et rapprochez-vous un peu.",
    "decoder_unavailable": "Le scanner n'a pas pu démarrer sur cet appareil. Saisissez l'identifiant de référence imprimé sous le code QR dans le champ ci-dessous.",
    "no_read_banner": "Impossible de lire ce code QR",
    "no_read_banner_sub": "Saisissez l'identifiant de référence imprimé sous le code QR dans le champ ci-dessous.",
    "no_read_hint": "Ce code QR ne peut pas être lu — il est peut-être flou, effacé ou de mauvaise qualité. Saisissez l'identifiant de référence imprimé sous le code QR dans le champ ci-dessous.",

    # ---- fsp_admin.html (connection state, counters, send) ----
    "current_program": "Programme",
    "status_online": "En ligne",
    "status_online_sub": "Vous avez Internet",
    "status_offline": "Pas de connexion",
    "status_offline_sub": "Vérifiez votre signal",
    "status_offline_mode": "Mode hors ligne",
    "status_offline_mode_sub": "L'application n'utilisera pas Internet",
    "work_offline": "Travailler hors ligne",
    "go_back_online": "Revenir en ligne",
    "beneficiaries_ready_to_scan": "Prêts à scanner :",
    "last_synced": "Dernière synchronisation :",
    "cannot_sync_offline": "Impossible de synchroniser hors ligne",
    "people_scanned": "Personnes scannées :",
    "total_amount": "Montant total :",
    "payments_ready_send": "Paiements prêts à envoyer",
    "last_payments_submit": "Dernier envoi :",
    "cannot_send_offline": "Impossible d'envoyer hors ligne",
    "payment_submit_success": "✅ Paiements envoyés avec succès !",
    "payment_submit_failed": "❌ Échec de l'envoi",
    "no_payments": "Aucun paiement prêt à envoyer.",

    # ---- beneficiary_offline.html ----
    "title": "Informations sur le bénéficiaire",
    "decryption_enabled": "Déchiffrement activé",
    "loading_photo": "Chargement de la photo…",
    "payment_amount": "Montant du paiement",
    "multiple_payments": "Ce bénéficiaire a plus d'un paiement enregistré. Sélectionnez celui que vous distribuez.",
    "select_payment": "Sélectionnez un paiement",
    "selected": "Sélectionné",
    "all_collected": "Tous les paiements pour ce bénéficiaire ont déjà été retirés.",
    "payment_approved": "Approuver",
    "payment_rejected": "Passer",
    "enter_value": "Saisir une valeur",

    # ---- success_offline.html ----
    "success_title": "Soumis avec succès",
    "success_message": "Vous pouvez maintenant scanner le bénéficiaire suivant.",
    "payments_ready": "Paiements prêts à envoyer à 121 :",
    "scan_next": "Scanner le bénéficiaire suivant",
    "go_home": "Terminé le scan ? Aller à la page d’accueil",

    # ---- invalid-qr.html ----
    "invalid_qr_title": "Code QR invalide",
    "invalid_qr_message": "Ce code QR ne peut pas être utilisé.",
    "checking_reason": "Vérification de la raison…",
    "reason_used": "Ce code QR a déjà été utilisé.",
    "reason_invalid": "Le code QR est invalide ou non reconnu.",
    "reason_no_record": "Aucun bénéficiaire correspondant n’a été trouvé.",
    "reason_database": "Problème de lecture des données hors ligne. Veuillez réessayer.",

    # ---- vouchers.html (upload + field order) ----
    "voucher_generator": "Générateur de bons",
    "csv_hint": "Téléchargez un fichier CSV ou Excel avec referenceId et d’autres champs à imprimer",
    "choose_csv": "Choisir un fichier CSV ou Excel…",
    "upload_csv": "Importer le fichier CSV ou Excel",
    "download_vouchers": "Télécharger les bons (PDF)",
    "choose_csv_alert": "Veuillez d'abord choisir un fichier CSV.",
    "upload_failed": "Échec du téléversement",
    "voucher_ready_singular": "bon prêt à télécharger",
    "voucher_ready_plural": "bons prêts à télécharger",
    "voucher_design_incomplete_title": "Conception du bon non terminée",
    "voucher_design_incomplete_desc": "Ce programme n'a pas encore de conception de bon. Définissez le titre, le sous-titre et les logos avant de générer les bons.",
    "voucher_design_go": "Concevoir ce bon",
    "voucher_fields_title": "Champs sur le bon",
    "voucher_fields_hint": "Faites glisser pour réordonner. Cliquez sur un nom pour modifier son affichage. Désactivez un champ pour l'exclure du bon. Enregistré pour ce programme.",
    "voucher_fields_none": "Aucune colonne imprimable n'a été trouvée dans ce fichier.",
    "voucher_fields_order_saved": "Ordre des champs enregistré",
    "voucher_fields_order_failed": "Impossible d'enregistrer l'ordre des champs",
    "voucher_fields_move_up": "Monter",
    "voucher_fields_move_down": "Descendre",
    "voucher_fields_sample": "ex.",
    "voucher_fields_rename": "Modifier le nom imprimé sur le bon",
    "voucher_fields_column": "Colonne de votre fichier",
    "voucher_fields_reset": "réinitialiser",
    "voucher_fields_reset_title": "Réinitialiser à",

    # ---- voucher_design.html ----
    "voucher_design_title": "Conception du bon",
    "voucher_design_settings": "Paramètres de conception",
    "voucher_design_settings_desc": "Ceux-ci remplacent les valeurs imprimées sur le bon. Chaque programme conserve sa propre conception.",
    "voucher_title_label": "Titre",
    "voucher_subtitle_label": "Sous-titre",
    "voucher_logos_label": "Logos",
    "voucher_left_logo": "Logo de gauche",
    "voucher_right_logo": "Logo de droite",
    "voucher_logo_upload": "Téléverser",
    "voucher_logo_remove": "supprimer",
    "voucher_no_logo": "Aucun logo",
    "voucher_logo_size": "Taille",
    "voucher_size_small": "Petit",
    "voucher_size_medium": "Moyen",
    "voucher_size_large": "Grand",
    "voucher_show_qr": "Afficher le code QR",
    "voucher_show_qr_hint": "Désactivez si les bons de ce programme n'ont pas besoin d'un code QR.",
    "voucher_live_preview": "Aperçu en direct",
    "voucher_preview_desc": "A5 paysage, correspondant à la mise en page du bon imprimé.",
    "voucher_save_design": "Enregistrer la conception",
    "voucher_design_saved": "Conception enregistrée",
    "voucher_design_save_failed": "Impossible d'enregistrer la conception",
    "lookup_min": "Saisissez au moins {n} caractères du code.",
    "lookup_many": "{n} correspondances possibles — continuez à saisir.",
    "lookup_one": "Une correspondance trouvée — appuyez sur Valider.",
    "lookup_none": "Aucune correspondance dans les enregistrements. Vérifiez le code.",
    "verify_only_note": "Vérification uniquement — votre compte ne peut pas enregistrer de paiements."
},

"ar": {
    # ---- Shared chrome (header / footer / buttons on every page) ----
    "footer_dev": "تم التطوير بواسطة 510 @ الصليب الأحمر الهولندي",
    "footer_support": "إذا كنت بحاجة إلى الدعم، تواصل مع jharrison@redcross.nl",
    "logout": "تسجيل الخروج",
    "back_to_dashboard": "عودة",
    "back_to_programs": "العودة إلى البرامج",
    "go_back": "عودة",
    "save": "حفظ",
    "saved_successfully": "تم الحفظ بنجاح",
    "failed_to_save": "فشل الحفظ",
    "remove": "إزالة",
    "submit": "إرسال",
    "program": "البرنامج",
    "program_id": "اختر البرنامج",
    "select_option": "اختر...",

    # ---- home.html ----
    "home_title": "مرحبًا",
    "home_question": "من أنت؟",
    "home_admin": "المسؤول",
    "home_fsp": "موظفو التوزيع",

    # ---- admin_login.html / fsp_login.html (errors raised in app.py) ----
    "login": "تسجيل دخول المسؤول",
    "fsp_login": "تسجيل دخول موظفي التوزيع",
    "enter_password": "أدخل كلمة المرور",
    "login_error": "اسم المستخدم أو كلمة المرور غير صحيحة.",
    "wrong_credentials": "اسم المستخدم أو كلمة المرور غير صحيحة.",
    "api_unreachable": "تعذّر الاتصال بالخادم. يرجى المحاولة مرة أخرى.",

    # ---- fsp_programs.html ----
    "program_select_title": "اختر برنامجاً",
    "program_select_subtitle": "اختر البرنامج الذي تريد العمل معه",

    # ---- admin_dashboard.html ----
    "admin_dashboard_title": "لوحة تحكم المسؤول",
    "admin_dashboard_subtitle": "إدارة إعدادات 121 Scan وأدوات القسائم",
    "config_system": "إعدادات النظام",
    "config_system_desc": "تعديل مفاتيح API ونقاط النهاية ومعلمات النظام.",
    "config_display": "تكوين الحقول المعروضة",
    "config_display_desc": "تكوين الحقول والمعلومات التي يتم عرضها لمقدمي الخدمات المالية.",
    "generate_vouchers": "إصدار القسائم",
    "generate_vouchers_desc": "إنشاء قسائم QR للطباعة والتوزيع.",
    "voucher_design": "تصميم القسيمة",
    "voucher_design_desc": "خصّص العنوان والعنوان الفرعي والشعارات المطبوعة على قسائم كل برنامج.",

    # ---- system_config.html ----
    "kobo_api_token": "رمز API",
    "kobo_token_hint": "يوجد في KoboToolbox ضمن إعدادات الحساب ← الأمان ← رمز API.",
    "verify_token": "التحقق من الرمز",
    "verifying": "جارٍ التحقق…",
    "verify_unreachable": "تعذّر الوصول إلى الخادم.",
    "enter_token_first": "أدخل رمزًا جديدًا أعلاه، ثم تحقق.",
    "token_invalid": "الرمز غير صالح لهذا الخادم.",
    "token_linked_to": "الرمز مرتبط بـ",
    "program_map_hint": "يرتبط كل برنامج 121 بنموذج KoboToolbox الذي يجمع بياناته.",

    # ---- config.html (field display configuration) ----
    "config_title": "تكوين الحقول المعروضة",
    "config_subtitle": "اختر الحقول التي تظهر أثناء المسح وحدد تسمياتها بكل لغة. يسحب التقاط الصور خياراته من نموذج KoboToolbox الخاص بالبرنامج.",
    "config_program_hint": "تنطبق الإعدادات أدناه على البرنامج المحدد",
    "config_select_or_delete": "يرجى تحديد حقل أو حذف الصفوف المميزة قبل الحفظ.",
    "fields_section_title": "الحقول المعروضة",
    "fields_section_hint": "اختر حقلاً، ثم اضبط تسميته لكل لغة. تُملأ التسميات تلقائيًا ويمكن تعديلها.",
    "field_key": "حقل الصورة في كوبا",
    "field_name_121": "اسم الحقل في 121",
    "label_en": "التسمية (EN)",
    "label_fr": "التسمية (FR)",
    "label_ar": "التسمية (AR)",
    "add_field": "إضافة حقل",
    "select_field_121": "اختر حقلاً",
    "editable": "قابل للتعديل",
    "editable_hint": "فعّل للسماح لموظفي التوزيع بتعديل هذا الحقل أثناء المسح. تُرسل التعديلات إلى 121 مع الدفعة.",
    "editable_on": "مُفعّل",
    "editable_off": "مُعطّل",
    "scan_config_title": "طريقة المسح",
    "scan_section_hint": "اختر كيفية تحديد هوية المستفيدين.",
    "enable_qr_scan": "تفعيل مسح رمز QR",
    "enable_qr_scan_hint": "عند الإيقاف، يستخدم هذا البرنامج إدخال الرمز يدويًا فقط (بدون كاميرا QR).",
    "photo_config_title": "إعدادات حقل الصورة",
    "photo_section_hint": "يرتبط بسؤال صورة في نموذج KoboToolbox الخاص بالبرنامج. تُملأ التسميات من Kobo.",
    "enable_photo_field": "تفعيل عرض حقل الصورة",
    "enable_photo_field_hint": "عند التفعيل، يمكن للموظفين إرفاق صورة بحقل الصورة المحدد في Kobo.",
    "kobo_prefill_note": "تم ملء التسميات من KoboToolbox — عدّلها إذا لزم الأمر",

    # ---- scan.html ----
    "scan_title": "مسح رمز QR",
    "scan_hint": "وجّه الكاميرا نحو رمز QR.",
    "code_entry_title": "إدخال الرمز",
    "code_entry_hint": "أدخل الرمز الفريد للمستفيد أدناه.",
    "no_qr_btn": "تعذّر مسح رمز QR؟",
    "manual_code_label": "أدخل الرمز يدويًا",
    "manual_code_placeholder": "اكتب أو الصق الرمز…",
    "manual_submit": "تأكيد",
    "start_camera": "بدء تشغيل الكاميرا",
    "waiting_to_start": "بانتظار البدء…",
    "requesting_camera": "جارٍ طلب تشغيل الكاميرا… إذا طُلِب منك ذلك، اضغط سماح.",
    "camera_denied": "تم رفض إذن الكاميرا أو أنها غير متاحة.",
    "scanning": "جارٍ المسح…",
    "starting_camera": "جارٍ بدء تشغيل الكاميرا…",
    "scan_trouble_hint": "لا تزال قراءة الرمز جارية — ثبّت الجهاز واقترب قليلاً.",
    "decoder_unavailable": "تعذّر تشغيل الماسح الضوئي على هذا الجهاز. أدخل المعرّف المرجعي المطبوع أسفل رمز QR في المربع أدناه.",
    "no_read_banner": "تعذّرت قراءة رمز QR هذا",
    "no_read_banner_sub": "اكتب المعرّف المرجعي المطبوع أسفل رمز QR في المربع أدناه.",
    "no_read_hint": "تعذّرت قراءة رمز QR هذا — قد يكون ضبابيًا أو باهتًا أو منخفض الجودة. أدخل المعرّف المرجعي المطبوع أسفل الرمز في المربع أدناه.",

    # ---- fsp_admin.html (connection state, counters, send) ----
    "current_program": "البرنامج",
    "status_online": "متصل",
    "status_online_sub": "لديك اتصال بالإنترنت",
    "status_offline": "لا يوجد اتصال",
    "status_offline_sub": "تحقق من الإشارة",
    "status_offline_mode": "وضع عدم الاتصال",
    "status_offline_mode_sub": "لن يستخدم التطبيق الإنترنت",
    "work_offline": "العمل دون اتصال",
    "go_back_online": "العودة إلى الاتصال",
    "beneficiaries_ready_to_scan": "جاهزون للمسح:",
    "last_synced": "آخر مزامنة:",
    "cannot_sync_offline": "لا يمكن المزامنة دون اتصال",
    "people_scanned": "الأشخاص الذين تم مسحهم:",
    "total_amount": "المبلغ الإجمالي:",
    "payments_ready_send": "المدفوعات الجاهزة للإرسال",
    "last_payments_submit": "آخر إرسال:",
    "cannot_send_offline": "لا يمكن الإرسال دون اتصال",
    "payment_submit_success": "✅ تم إرسال المدفوعات بنجاح!",
    "payment_submit_failed": "❌ فشل الإرسال",
    "no_payments": "لا توجد مدفوعات جاهزة للإرسال.",

    # ---- beneficiary_offline.html ----
    "title": "معلومات المستفيد",
    "decryption_enabled": "فك التشفير مُفعّل",
    "loading_photo": "جارٍ تحميل الصورة…",
    "payment_amount": "مبلغ الدفع",
    "multiple_payments": "لدى هذا المستفيد أكثر من دفعة مسجّلة. اختر الدفعة التي توزّعها.",
    "select_payment": "اختر دفعة",
    "selected": "محدد",
    "all_collected": "تم بالفعل استلام جميع المدفوعات لهذا المستفيد.",
    "payment_approved": "موافقة",
    "payment_rejected": "تخطي",
    "enter_value": "أدخل القيمة",

    # ---- success_offline.html ----
    "success_title": "تم الإرسال بنجاح",
    "success_message": "يمكنك الآن مسح المستفيد التالي.",
    "payments_ready": "المدفوعات الجاهزة للإرسال إلى 121:",
    "scan_next": "مسح المستفيد التالي",
    "go_home": "هل انتهيت من المسح؟ اذهب إلى الصفحة الرئيسية",

    # ---- invalid-qr.html ----
    "invalid_qr_title": "رمز QR غير صالح",
    "invalid_qr_message": "لا يمكن استخدام رمز QR هذا.",
    "checking_reason": "جارٍ التحقق من السبب…",
    "reason_used": "تم استخدام رمز QR هذا سابقًا.",
    "reason_invalid": "رمز QR غير صالح أو غير معروف.",
    "reason_no_record": "لا يوجد أي مستفيد مطابق لهذا الرمز.",
    "reason_database": "حدثت مشكلة في قراءة البيانات دون اتصال. حاول مرة أخرى.",

    # ---- vouchers.html (upload + field order) ----
    "voucher_generator": "مولّد القسائم",
    "csv_hint": "قم بتحميل ملف CSV أو Excel يحتوي على referenceId وأي حقول إضافية للطباعة",
    "choose_csv": "اختر ملف CSV أو Excel…",
    "upload_csv": "تحميل ملف CSV أو Excel",
    "download_vouchers": "تنزيل القسائم (PDF)",
    "choose_csv_alert": "يرجى اختيار ملف CSV أولاً.",
    "upload_failed": "فشل الرفع",
    "voucher_ready_singular": "قسيمة جاهزة للتنزيل",
    "voucher_ready_plural": "قسائم جاهزة للتنزيل",
    "voucher_design_incomplete_title": "لم يكتمل تصميم القسيمة",
    "voucher_design_incomplete_desc": "لا يحتوي هذا البرنامج على تصميم قسيمة بعد. حدّد العنوان والعنوان الفرعي والشعارات قبل إنشاء القسائم.",
    "voucher_design_go": "تصميم هذه القسيمة",
    "voucher_fields_title": "الحقول على القسيمة",
    "voucher_fields_hint": "اسحب لإعادة الترتيب. انقر على الاسم لتغيير طريقة طباعته. أوقف أي حقل لاستثنائه من القسيمة. يُحفظ ذلك لهذا البرنامج.",
    "voucher_fields_none": "لم يتم العثور على أعمدة قابلة للطباعة في هذا الملف.",
    "voucher_fields_order_saved": "تم حفظ ترتيب الحقول",
    "voucher_fields_order_failed": "تعذّر حفظ ترتيب الحقول",
    "voucher_fields_move_up": "أعلى",
    "voucher_fields_move_down": "أسفل",
    "voucher_fields_sample": "مثال:",
    "voucher_fields_rename": "تعديل الاسم المطبوع على القسيمة",
    "voucher_fields_column": "العمود في ملفك",
    "voucher_fields_reset": "إعادة تعيين",
    "voucher_fields_reset_title": "إعادة التعيين إلى",

    # ---- voucher_design.html ----
    "voucher_design_title": "تصميم القسيمة",
    "voucher_design_settings": "إعدادات التصميم",
    "voucher_design_settings_desc": "تحل هذه محل القيم المطبوعة على القسيمة. يحتفظ كل برنامج بتصميمه الخاص.",
    "voucher_title_label": "العنوان",
    "voucher_subtitle_label": "العنوان الفرعي",
    "voucher_logos_label": "الشعارات",
    "voucher_left_logo": "الشعار الأيسر",
    "voucher_right_logo": "الشعار الأيمن",
    "voucher_logo_upload": "رفع",
    "voucher_logo_remove": "إزالة",
    "voucher_no_logo": "لا يوجد شعار",
    "voucher_logo_size": "الحجم",
    "voucher_size_small": "صغير",
    "voucher_size_medium": "متوسط",
    "voucher_size_large": "كبير",
    "voucher_show_qr": "إظهار رمز QR",
    "voucher_show_qr_hint": "أوقف التشغيل إذا كانت قسائم هذا البرنامج لا تحتاج إلى رمز QR.",
    "voucher_live_preview": "معاينة مباشرة",
    "voucher_preview_desc": "A5 أفقي، مطابق لتخطيط القسيمة المطبوعة.",
    "voucher_save_design": "حفظ التصميم",
    "voucher_design_saved": "تم حفظ التصميم",
    "voucher_design_save_failed": "تعذّر حفظ التصميم",
    "lookup_min": "أدخل {n} أحرف على الأقل من الرمز.",
    "lookup_many": "{n} تطابقات محتملة — تابع الإدخال.",
    "lookup_one": "تم العثور على تطابق واحد — اضغط تأكيد.",
    "lookup_none": "لا يوجد تطابق في السجلات المحفوظة. تحقق من الرمز.",
    "verify_only_note": "التحقق فقط — لا يمكن لحسابك تسجيل المدفوعات."
},

}


@app.route("/instance-static/<filename>")
def instance_static(filename):
    context = os.getenv("SCANDROID_CONTEXT", "local")
    azure_path = f"/home/site/configs/{context}/static"

    # Azure instance-specific logos
    instance_file = os.path.join(azure_path, filename)
    if os.path.exists(instance_file):
        resp = send_from_directory(azure_path, filename)
    else:
        # Local fallback for development
        resp = send_from_directory("static", filename)

    # These files (e.g. voucher logos) are overwritten in place on re-upload,
    # so the browser must revalidate rather than serve a cached copy. The URL
    # also carries a ?v=<mtime> cache-buster, but we set no-cache here as a
    # second line of defence against stale logos after an upload.
    resp.headers["Cache-Control"] = "no-cache, must-revalidate"
    return resp


@app.route("/")
def landing_page():
    lang = request.args.get("lang", "en")
    return render_template(
        "home.html", lang=lang, t=translations.get(lang, translations["en"])
    )


@app.route("/admin-login", methods=["GET", "POST"])
def admin_login():
    # language handling
    lang = (
        request.args.get("lang")
        or request.form.get("lang")
        or session.get("lang", "en")
    )
    session["lang"] = lang
    t = translations.get(lang, translations["en"])

    # GET → show login page
    if request.method == "GET":
        return render_template("admin_login.html", lang=lang, t=t, error=None)

    # POST → authenticate against 121 API using system_config.json
    username = request.form.get("username")
    password = request.form.get("password")

    config = load_config()
    base_url = config.get("url121")
    if not base_url:
        return render_template(
            "admin_login.html",
            lang=lang,
            t=t,
            error="❌ Missing url121 in system configuration.",
        )

    login_url = f"{base_url}/api/users/login"
    login_payload = {"username": username, "password": password}

    try:
        res = requests.post(login_url, json=login_payload, timeout=10)

    except Exception:
        # API unreachable
        return render_template(
            "admin_login.html",
            lang=lang,
            t=t,
            error=t.get("api_unreachable", "Unable to reach login server."),
        )

    # ✔ Success
    if res.status_code == 201:
        _store_login_session("admin", username, res)
        return redirect(url_for("admin_dashboard", lang=lang))

    # ❌ Wrong username/password
    if res.status_code in (400, 401):
        return render_template(
            "admin_login.html",
            lang=lang,
            t=t,
            error=t.get("wrong_credentials", "Incorrect username or password."),
        )

    # ❌ Any other response
    return render_template(
        "admin_login.html", lang=lang, t=t, error=f"Login failed ({res.status_code})."
    )


@app.route("/admin-dashboard")
def admin_dashboard():
    if not session.get("admin_logged_in"):
        return redirect(url_for("admin_login", lang=request.args.get("lang", "en")))

    lang = request.args.get("lang", "en")
    t = translations.get(lang, translations["en"])

    return render_template(
        "admin_dashboard.html", lang=lang, t=t, username=session.get("admin_username")
    )


@app.route("/admin-logout")
def admin_logout():
    lang = request.args.get("lang", "en")
    _clear_login_session("admin")
    return redirect(url_for("admin_login", lang=lang))


from flask import request, session, redirect, url_for, flash, render_template, jsonify
import requests
from requests.auth import HTTPBasicAuth
import json


@app.route("/system-config", methods=["GET", "POST"])
def system_config():
    if not session.get("admin_logged_in"):
        return redirect(url_for("admin_login", lang=request.args.get("lang", "en")))

    lang = request.args.get("lang", "en")
    t = translations.get(lang, translations["en"])
    config = load_config()

    # ============================================================
    # POST: SAVE CONFIG
    # ============================================================
    if request.method == "POST":
        updated = config.copy()

        # ---- Global fields ----
        # Always save these from the form
        for key in ["KOBO_SERVER", "url121"]:
            updated[key] = request.form.get(key, "").strip()

        # Only overwrite if a new value was submitted
        for key in ["KOBO_TOKEN", "username121", "password121", "ENCRYPTION_KEY"]:
            submitted = request.form.get(key, "").strip()
            if submitted:
                updated[key] = submitted
            # else: keep existing value from config.copy()

        # ---- Program mappings ----
        programs = []
        program_ids = request.form.getlist("PROGRAMS[][programId]")
        asset_ids = request.form.getlist("PROGRAMS[][koboAssetId]")

        kobo_server = updated.get("KOBO_SERVER")
        kobo_token = updated.get("KOBO_TOKEN")

        for pid, asset_id in zip(program_ids, asset_ids):
            if not pid or not asset_id:
                continue

            entry = {"programId": int(pid), "koboAssetId": asset_id.strip()}

            # ---- Validate Kobo asset ----
            try:
                r = requests.get(
                    f"{kobo_server}/api/v2/assets/{asset_id}/?format=json",
                    headers={"Authorization": f"Token {kobo_token}"},
                    timeout=10,
                )
                if r.status_code == 200:
                    j = r.json()
                    entry["koboFormName"] = j.get("name")
                    entry["koboFormOwner"] = j.get("owner__username")
            except Exception as e:
                print("Kobo validation error:", e)

            programs.append(entry)

        updated["PROGRAMS"] = programs

        save_config(updated)
        flash(t["saved_successfully"])
        return redirect(url_for("system_config", lang=lang))

    # ============================================================
    # GET: LOAD DATA FOR UI
    # ============================================================
    url121 = config.get("url121")
    username121 = config.get("username121")
    password121 = config.get("password121")

    token = None
    program_ids = []
    program_options = []

    # ---- Login to 121 ----
    if url121 and username121 and password121:
        try:
            r = requests.post(
                f"{url121}/api/users/login",
                json={"username": username121, "password": password121},
            )
            if r.status_code == 201:
                j = r.json()
                token = j.get("access_token_general")
                program_ids = [int(pid) for pid in j.get("permissions", {}).keys()]
        except Exception as e:
            print("121 login failed:", e)

    # ---- Load program titles ----
    if token:
        for pid in program_ids:
            try:
                r = requests.get(
                    f"{url121}/api/programs/{pid}",
                    cookies={"access_token_general": token},
                )
                if r.status_code == 200:
                    pdata = r.json()
                    titles = pdata.get("titlePortal", {})
                    title = titles.get(lang) or next(
                        iter(titles.values()), f"Program {pid}"
                    )
                    program_options.append({"id": pid, "title": title})
            except Exception as e:
                print(f"Program load failed ({pid}):", e)

    # Existing mappings for UI
    program_mappings = config.get("PROGRAMS", [])

    return render_template(
        "system_config.html",
        config=config,
        program_options=program_options,
        program_mappings=program_mappings,
        username=session.get("admin_username"),
        lang=lang,
        t=t,
    )

@app.route("/verify-kobo-token", methods=["POST"])
def verify_kobo_token():
    if not session.get("admin_logged_in"):
        return jsonify({"ok": False, "error": "Unauthorized"}), 401

    data = request.get_json(silent=True) or {}
    server = (data.get("server") or "").strip().rstrip("/")
    token = (data.get("token") or "").strip()

    if not server or not token:
        return jsonify({"ok": False, "error": "Server and token are required."}), 400

    headers = {"Authorization": f"Token {token}"}

    for path in ("/api/v2/me/?format=json", "/me/?format=json"):
        try:
            r = requests.get(f"{server}{path}", headers=headers, timeout=10)
        except Exception as e:
            print("Kobo verify failed:", e)
            return jsonify({"ok": False, "error": "Could not reach the server."}), 200

        if r.status_code == 200:
            try:
                j = r.json()
            except Exception:
                continue
            return jsonify({
                "ok": True,
                "username": j.get("username", ""),
                "email": j.get("email", ""),
            })

        if r.status_code in (401, 403):
            return jsonify({"ok": False, "error": "Token not valid for this server."}), 200

    return jsonify({"ok": False, "error": "Token not valid for this server."}), 200

@app.route("/api/program-attributes/<int:program_id>")
def api_program_attributes(program_id):
    if not session.get("admin_logged_in"):
        return jsonify({"error": "Unauthorized"}), 401

    system_config = load_config()
    url121 = system_config.get("url121")

    if not url121:
        return jsonify({"attributes": [], "kobo_image_fields": []})

    attributes = []
    kobo_image_fields = []

    try:
        login_resp = requests.post(
            f"{url121}/api/users/login",
            json={
                "username": system_config.get("username121", ""),
                "password": system_config.get("password121", ""),
            },
            timeout=10,
        )

        if login_resp.status_code == 201:
            token = login_resp.json().get("access_token_general")
            cookies = {"access_token_general": token}

            # Registration attributes
            r = requests.get(
                f"{url121}/api/programs/{program_id}", cookies=cookies, timeout=10
            )
            if r.status_code == 200:
                for attr in r.json().get("programRegistrationAttributes", []):
                    name = attr.get("name")
                    if not name:
                        continue
                    labels = attr.get("label") or {}
                    label = labels.get("en") or next(iter(labels.values()), name)
                    attributes.append({"name": name, "label": label})

    except Exception as e:
        print(f"[api_program_attributes] Error: {e}")

    # Kobo image fields for this program
    try:
        programs = system_config.get("PROGRAMS", [])
        program = next(
            (p for p in programs if str(p.get("programId")) == str(program_id)), None
        )
        if program:
            asset_id = program.get("koboAssetId")
            kobo_token = system_config.get("KOBO_TOKEN")
            kobo_server = system_config.get("KOBO_SERVER", "https://kobo.ifrc.org")

            if asset_id and kobo_token:
                r = requests.get(
                    f"{kobo_server}/api/v2/assets/{asset_id}/?format=json",
                    headers={"Authorization": f"Token {kobo_token}"},
                    timeout=10,
                )
                if r.status_code == 200:
                    survey = r.json().get("content", {}).get("survey", [])
                    for item in survey:
                        if not isinstance(item, dict) or item.get("type") != "image":
                            continue
                        raw_label = item.get("label")
                        if isinstance(raw_label, list) and raw_label:
                            label = (
                                raw_label[0]
                                if isinstance(raw_label[0], str)
                                else next(iter(raw_label[0].values()), "")
                            )
                        elif isinstance(raw_label, dict):
                            label = raw_label.get("en") or next(
                                iter(raw_label.values()), ""
                            )
                        else:
                            label = str(raw_label or "")
                        xpath = (
                            item.get("$xpath", "")
                            .replace("/data/", "")
                            .replace("data/", "")
                            .strip("/")
                        )
                        name = xpath or item.get("name", "")
                        if name:
                            kobo_image_fields.append({"name": name, "label": label})
    except Exception as e:
        print(f"[api_program_attributes] Kobo error: {e}")

    return jsonify({"attributes": attributes, "kobo_image_fields": kobo_image_fields})


@app.route("/config", methods=["GET", "POST"])
def config_page():
    if not session.get("admin_logged_in"):
        return redirect(url_for("admin_login", lang=request.args.get("lang", "en")))

    lang = request.args.get("lang", "en")
    username = session.get("admin_username", "Admin")

    # ------------------------------------------------------
    # POST (SAVE DISPLAY CONFIG – PER PROGRAM)
    # ------------------------------------------------------
    if request.method == "POST":
        config_data = request.get_json()

        program_id = str(config_data.pop("programId"))
        config_data.pop("COLUMN_TO_MATCH", None)

        try:
            full_config = load_display_config()
            if "programs" not in full_config:
                full_config["programs"] = {}

            full_config["programs"][program_id] = config_data
            save_display_config(full_config)

            return jsonify({"success": True})
        except Exception as e:
            return jsonify({"success": False, "error": str(e)}), 500

    # ------------------------------------------------------
    # GET (LOAD PAGE)
    # ------------------------------------------------------

    # Load display config
    try:
        config_data = load_display_config()
    except Exception:
        config_data = {}

    # Load system config
    try:
        system_config = load_config()
    except Exception:
        system_config = {}

    programs_raw = system_config.get("PROGRAMS", [])
    programs = []  # UI list (id + title)
    program_lookup = {}  # Logic lookup (full object)

    url121 = system_config.get("url121")

    # ------------------------------------------------------
    # Resolve program titles from 121
    # ------------------------------------------------------
    if url121 and programs_raw:
        try:
            login_resp = requests.post(
                f"{url121}/api/users/login",
                json={
                    "username": system_config.get("username121", ""),
                    "password": system_config.get("password121", ""),
                },
            )

            if login_resp.status_code == 201:
                token = login_resp.json().get("access_token_general")
                cookies = {"access_token_general": token}

                for p in programs_raw:
                    pid = p.get("programId")
                    title = str(pid)

                    try:
                        r = requests.get(
                            f"{url121}/api/programs/{pid}", cookies=cookies, timeout=10
                        )
                        if r.status_code == 200:
                            titles = r.json().get("titlePortal", {})
                            title = titles.get("en") or next(
                                iter(titles.values()), title
                            )
                    except Exception:
                        pass

                    programs.append({"id": str(pid), "title": title})
                    program_lookup[str(pid)] = p

        except Exception as e:
            print("Program title lookup failed:", e)

    # Safety fallback
    if not programs:
        for p in programs_raw:
            pid = str(p.get("programId"))
            programs.append({"id": pid, "title": pid})
            program_lookup[pid] = p

    # ------------------------------------------------------
    # Active program selection
    # ------------------------------------------------------
    active_program_id = request.args.get("programId")
    if not active_program_id and programs:
        active_program_id = programs[0]["id"]
    active_program_id = str(active_program_id) if active_program_id else None

    # ------------------------------------------------------
    # Backward compatibility (single → multi)
    # ------------------------------------------------------
    if "programs" not in config_data:
        default_program_id = active_program_id or "default"
        config_data = {
            "programs": {
                default_program_id: {
                    "fields": config_data.get("fields", []),
                    "photo": config_data.get(
                        "photo",
                        {
                            "enabled": True,
                            "labels": {"en": "Photo", "fr": "Photo", "ar": "صورة"},
                        },
                    ),
                }
            }
        }

    # ------------------------------------------------------
    # Resolve active program Kobo + 121 config
    # ------------------------------------------------------
    program_id = None
    asset_id = None

    for prog in system_config.get("PROGRAMS", []):
        if str(prog.get("programId")) == str(active_program_id):
            program_id = prog.get("programId")
            asset_id = prog.get("koboAssetId")
            break

    program_title = system_config.get("programTitle", "")
    column_to_match_121 = None
    allowed_attributes = []
    kobo_image_fields = []

    # ------------------------------------------------------
    # 121 PROGRAM ATTRIBUTES
    # ------------------------------------------------------
    if url121 and program_id:
        try:
            login_resp = requests.post(
                f"{url121}/api/users/login",
                json={
                    "username": system_config.get("username121", ""),
                    "password": system_config.get("password121", ""),
                },
            )

            if login_resp.status_code == 201:
                token = login_resp.json().get("access_token_general")
                cookies = {"access_token_general": token}

                # Column to match
                # The inner `break` left the outer loop running, so this page
                # took the LAST configuration while offline_sync.py took the
                # FIRST — they actively disagreed on any multi-FSP program, and
                # this page then persisted its answer into
                # COLUMN_TO_MATCH_PER_PROGRAM where the sync would read it back.
                # Only report a value when every configuration agrees.
                try:
                    _map = get_fsp_column_map(program_id)
                    _unique = sorted(set(_map.values()))
                    if len(_unique) == 1:
                        column_to_match_121 = _unique[0]
                    elif len(_unique) > 1:
                        column_to_match_121 = None
                        print(
                            f"[config] program {program_id} has per-FSP match "
                            f"columns {_map} — resolved per payment at sync time."
                        )
                except Exception as e:
                    print(f"[config] fsp-configurations lookup failed: {e}")

                # Registration attributes
                try:
                    r = requests.get(
                        f"{url121}/api/programs/{program_id}",
                        cookies=cookies,
                        timeout=10,
                    )
                    if r.status_code == 200:
                        for attr in r.json().get("programRegistrationAttributes", []):
                            name = attr.get("name")
                            if not name:
                                continue

                            labels = attr.get("label") or {}
                            label = labels.get("en") or next(
                                iter(labels.values()), name
                            )

                            allowed_attributes.append({"name": name, "label": label})
                except Exception:
                    pass

        except Exception as e:
            print("121 lookup failed:", e)

    # ------------------------------------------------------
    # Strict mode – drop invalid fields per program
    # ------------------------------------------------------
    if allowed_attributes and active_program_id:
        allowed_names = {a["name"] for a in allowed_attributes}
        pdata = config_data.get("programs", {}).get(active_program_id)
        if pdata:
            pdata["fields"] = [
                f for f in pdata.get("fields", []) if f.get("key") in allowed_names
            ]

            # Same treatment for the lookup field. A configured attribute that
            # 121 no longer exposes can never match anything typed into the
            # scan box, and would fail silently in the field — so it reverts to
            # the reference ID, which always exists.
            _lookup = pdata.get("lookup") or {}
            _lookup_key = _lookup.get("field")
            if _lookup_key and _lookup_key != "referenceId" \
                    and _lookup_key not in allowed_names:
                print(
                    f"[config] program {active_program_id}: lookup field "
                    f"'{_lookup_key}' is not a 121 attribute — reverting to "
                    "referenceId."
                )
                pdata["lookup"] = {"field": "referenceId"}

    # ------------------------------------------------------
    # Kobo image fields
    # ------------------------------------------------------
    def get_full_kobo_path(item):
        if not isinstance(item, dict):
            return None

        if "$xpath" in item:
            return item["$xpath"].replace("/data/", "").replace("data/", "").strip("/")

        parts = []
        current = item
        while isinstance(current, dict):
            if current.get("name"):
                parts.append(current["name"])
            current = current.get("parent")

        return "/".join(reversed(parts)) if parts else None

    try:
        token = system_config.get("KOBO_TOKEN")
        kobo_server = system_config.get("KOBO_SERVER", "https://kobo.ifrc.org")

        if token and asset_id:
            r = requests.get(
                f"{kobo_server}/api/v2/assets/{asset_id}/?format=json",
                headers={"Authorization": f"Token {token}"},
                timeout=10,
            )

            if r.status_code == 200:
                survey = r.json().get("content", {}).get("survey", [])
                for item in survey:
                    if not isinstance(item, dict):
                        continue

                    if item.get("type") == "image":
                        raw_label = item.get("label")

                        if isinstance(raw_label, dict):
                            label = (
                                raw_label.get("English")
                                or raw_label.get("en")
                                or next(
                                    iter(raw_label.values()), get_full_kobo_path(item)
                                )
                            )

                        elif isinstance(raw_label, list) and raw_label:
                            first = raw_label[0]
                            if isinstance(first, dict):
                                label = (
                                    first.get("English")
                                    or first.get("en")
                                    or next(
                                        iter(first.values()), get_full_kobo_path(item)
                                    )
                                )
                            else:
                                label = str(first)

                        else:
                            label = get_full_kobo_path(item)

                        path = get_full_kobo_path(item)
                        if path:
                            kobo_image_fields.append({"name": path, "label": label})
    except Exception as e:
        print("Kobo lookup failed:", e)

    # ------------------------------------------------------
    # Strict mode – image field per program
    # ------------------------------------------------------
    if kobo_image_fields:
        valid = {i["name"] for i in kobo_image_fields}
        for pdata in config_data.get("programs", {}).values():
            if pdata.get("photo", {}).get("field_name") not in valid:
                pdata.setdefault("photo", {})["field_name"] = ""

    # COLUMN_TO_MATCH_PER_PROGRAM is no longer written or read. The match column
    # is resolved per payment from the FSP configuration at sync and submit
    # time (INVARIANT 5: no program-scoped copy can go stale, because none
    # exists). Existing entries in system_config.json are inert and can be
    # deleted.

    # ------------------------------------------------------
    # RENDER
    # ------------------------------------------------------
    return render_template(
        "config.html",
        full_config=config_data,
        programs=programs,
        active_program_id=active_program_id,
        system_config=system_config,
        allowed_attributes=allowed_attributes,
        kobo_image_fields=kobo_image_fields,
        # No global fallback: when the program's FSP configurations disagree the
        # honest answer is "resolved per payment", not one arbitrary column.
        column_to_match_121=column_to_match_121,
        lang=lang,
        username=username,
        program_title=program_title,
        t=translations.get(lang, translations["en"]),
    )


@app.route("/logout")
def logout():
    lang = request.args.get("lang", "en")
    session.clear()
    return redirect(url_for("login", lang=lang))


@app.route("/fsp-login", methods=["GET", "POST"])
def fsp_login():
    lang = request.args.get("lang", "en")
    t = translations.get(lang, translations["en"])
    error = None

    config = load_config()
    base_url = config.get("url121")
    if not base_url:
        return render_template(
            "fsp_login.html", lang=lang, t=t, error="❌ Missing url121"
        )

    login_url = f"{base_url}/api/users/login"

    if request.method == "POST":
        username = request.form.get("username")
        password = request.form.get("password")

        try:
            res = requests.post(
                login_url, json={"username": username, "password": password}, timeout=8
            )

            if res.status_code == 201:
                _store_login_session("fsp", username, res)
                return redirect(url_for("fsp_program_selector"))

            elif res.status_code in (400, 401):
                error = t["login_error"]

            else:
                error = f"Login failed ({res.status_code})."

        except Exception:
            error = t["login_error"]

    return render_template("fsp_login.html", lang=lang, t=t, error=error)


@app.route("/fsp-programs")
def fsp_program_selector():
    # ------------------------------------------------------
    # Auth guard
    # ------------------------------------------------------
    if not session.get("fsp_logged_in"):
        return redirect(url_for("fsp_login"))

    # ------------------------------------------------------
    # Language + translations (MISSING BEFORE)
    # ------------------------------------------------------
    lang = request.args.get("lang", "en")
    t = translations.get(lang, translations["en"])
    username = session.get("fsp_username")

    # ------------------------------------------------------
    # Always initialise programs (CRITICAL FIX)
    # ------------------------------------------------------
    programs = []

    # ------------------------------------------------------
    # Load system config
    # ------------------------------------------------------
    try:
        system_config = load_config()
    except Exception:
        system_config = {}

    programs_raw = system_config.get("PROGRAMS", [])
    url121 = system_config.get("url121")

    # ------------------------------------------------------
    # Resolve program titles from 121
    # ------------------------------------------------------
    if url121 and programs_raw:
        try:
            login_resp = requests.post(
                f"{url121}/api/users/login",
                json={
                    "username": system_config.get("username121", ""),
                    "password": system_config.get("password121", ""),
                },
                timeout=10,
            )

            if login_resp.status_code == 201:
                token = login_resp.json().get("access_token_general")
                cookies = {"access_token_general": token}

                for p in programs_raw:
                    pid = p.get("programId")
                    title = str(pid)

                    try:
                        r = requests.get(
                            f"{url121}/api/programs/{pid}", cookies=cookies, timeout=10
                        )
                        if r.status_code == 200:
                            titles = r.json().get("titlePortal", {})
                            title = titles.get("en") or next(
                                iter(titles.values()), title
                            )
                    except Exception:
                        pass

                    programs.append({"id": str(pid), "title": title})

        except Exception as e:
            print("Program lookup failed:", e)

    # ------------------------------------------------------
    # Fallback: titles = program IDs
    # ------------------------------------------------------
    if not programs:
        programs = [
            {"id": str(p.get("programId")), "title": str(p.get("programId"))}
            for p in programs_raw
        ]

    # ------------------------------------------------------
    # Render selector page
    # ------------------------------------------------------
    return render_template(
        "fsp_programs.html",
        programs=programs,
        lang=lang,
        t=t,
        username=username,
        program_title="Program Selector",
    )


@app.route("/select-program/<program_id>")
def select_program(program_id):
    if not session.get("fsp_logged_in"):
        return redirect(url_for("fsp_login"))

    # keep language if present
    lang = request.args.get("lang", "en")

    # store selection in session
    session["fsp_program_id"] = str(program_id)

    # redirect into the admin page WITH query param
    return redirect(url_for("fsp_admin", program_id=str(program_id), lang=lang))


@app.route("/fsp-admin")
def fsp_admin():
    # --- auth ---
    if not session.get("fsp_logged_in"):
        return redirect(url_for("fsp_login"))

    # --- language ---
    lang = request.args.get("lang", "en")
    t = translations.get(lang, translations["en"])

    # --- program resolution (URL first, then session) ---
    program_id = request.args.get("program_id") or session.get("fsp_program_id")
    if not program_id:
        return redirect(url_for("fsp_program_selector", lang=lang))

    program_id = str(program_id)

    # --- load configs ---
    system_config = load_config()
    _full_display = load_display_config()
    # Pass full display config to template so it can be stored in IndexedDB
    display_config = _full_display

    # Kept only to satisfy the template's COLUMN_TO_MATCH variable. It is no
    # longer used for reconciliation — the device reads the column off each
    # tranche. Populated only when every FSP configuration agrees.
    _cols = set(get_fsp_column_map(program_id).values())
    column_to_match = _cols.pop() if len(_cols) == 1 else ""
    programs = system_config.get("PROGRAMS", [])

    # --- resolve program ---
    program = next((p for p in programs if str(p.get("programId")) == program_id), None)

    if not program:
        return redirect(url_for("fsp_program_selector", lang=lang))

    # --- fetch 121 program title ---
    program_title = f"Program {program_id}"
    url121 = system_config.get("url121")
    if url121:
        try:
            login_resp = requests.post(
                f"{url121}/api/users/login",
                json={
                    "username": system_config.get("username121", ""),
                    "password": system_config.get("password121", ""),
                },
                timeout=8,
            )
            if login_resp.status_code == 201:
                token = login_resp.json().get("access_token_general")
                r = requests.get(
                    f"{url121}/api/programs/{program_id}",
                    cookies={"access_token_general": token},
                    timeout=8,
                )
                if r.status_code == 200:
                    titles = r.json().get("titlePortal", {})
                    program_title = titles.get(lang) or next(
                        iter(titles.values()), program_title
                    )
        except Exception as e:
            print(f"[fsp_admin] Failed to fetch 121 program title: {e}")

    username = session.get("fsp_username")

    # --- IMPORTANT: persist program for later routes ---
    session["fsp_program_id"] = program_id

    # Two separate questions, because they fail at two different endpoints.
    #
    # Tri-state (see _has_program_permission): only a DEFINITE False hides
    # anything. None means 121's login response carried no readable permissions
    # map — we fail open there, exactly as /submit-payments does, because
    # blocking every distribution over a changed response shape would be far
    # worse than a late 403 from 121 itself.

    # 1. Can this operator reconcile payments? Gates the Send block.
    can_update_payments = (
        _has_program_permission("fsp", program_id, "payment.update") is not False
    )

    # 2. Can this operator update registrations? Gates the editable fields,
    #    which are pushed by PATCH /api/programs/{id}/registrations. Without
    #    these, 121 rejects the whole bulk update AFTER the payments have
    #    already gone through — so the input must never appear in the first
    #    place.
    can_update_registrations = all(
        _has_program_permission("fsp", program_id, perm) is not False
        for perm in REGISTRATION_UPDATE_PERMISSIONS
    )

    # An edit can only reach 121 via Send, so it needs both. Showing an input
    # to someone who cannot Send would collect data into a dead end.
    can_edit_fields = can_update_payments and can_update_registrations

    return render_template(
        "fsp_admin.html",
        COLUMN_TO_MATCH=column_to_match,
        display_config=display_config,
        lang=lang,
        t=t,
        config=system_config,
        program_title=program_title,
        program_id=program_id,  # ✅ this feeds ACTIVE_PROGRAM_ID in JS
        username=username,
        can_update_payments=can_update_payments,
        can_edit_fields=can_edit_fields,
    )

# ---------------------------------------------------------------------------
# BACKGROUND SYNC JOB RUNNER
#
# offline_sync.py can now run for many minutes (it pulls every waiting
# transaction, registration and photo — a 1,000-registration program is minutes,
# not seconds). Running it inside the request, as this used to, blocks the
# gunicorn worker for that whole time. On Azure that means:
#   - /ping stops answering, so fsp_admin.html's reachability probe fails three
#     times and the page declares itself OFFLINE mid-sync;
#   - the Azure front-end kills any request idle for 230s, so a big sync would
#     return 502 even on a healthy worker.
#
# So /sync-fsp now STARTS the sync in a background thread and returns
# immediately. Progress is parsed out of offline_sync.py's stdout and written to
# a small job file, which /sync-status serves to the page.
#
# The job file (not a module global) is deliberate: gunicorn recycles workers,
# and the status poll can land on a different worker than the one that started
# the job. A file survives both.
# ---------------------------------------------------------------------------
import threading
import time
import uuid as _uuid_mod
import re as _sync_re
from collections import deque as _sync_deque

SYNC_JOB_DIR = "offline-cache"
SYNC_JOB_FILE = os.path.join(SYNC_JOB_DIR, "sync_job.json")

# If the heartbeat is older than this, the worker that owned the job is gone
# (recycled, crashed, container restarted). Report it rather than showing a
# spinner forever.
SYNC_STALE_SECONDS = 90

_sync_job_lock = threading.Lock()

# Progress markers emitted by offline_sync.py. Keep these in step with the
# logger.info() calls in that file — if a message is reworded, progress silently
# stops advancing (the sync itself is unaffected).
_RE_OPEN_PAYMENTS = _sync_re.compile(r"\[INFO\]\s+(\d+)\s+open payment\(s\) selected")
_RE_PAYMENT_DONE = _sync_re.compile(r"\[INFO\]\s+paymentId=(\S+):\s+(\d+)\s+transaction")
_RE_BENEFICIARIES = _sync_re.compile(r"\[INFO\]\s+(\d+)\s+beneficiar\(y/ies\) with open payments")
_RE_REGISTRATIONS = _sync_re.compile(r"\[INFO\]\s+Registrations:\s+(\d+)\s+requested,\s+(\d+)\s+fetched,\s+(\d+)\s+failed")
_RE_PHOTO_OK = _sync_re.compile(r"^\[OK\]\s+Photo downloaded & encrypted")
_RE_FINAL = _sync_re.compile(r"^(\d+)\s+beneficiaries ready for offline validation")
# Fine-grained counters emitted by offline_sync.py during the two long phases.
# These are what make the progress bar move continuously rather than jumping
# between phase boundaries.
_RE_PROGRESS = _sync_re.compile(r"\[PROGRESS\]\s+(registrations|photos)\s+(\d+)/(\d+)")


def _sync_job_read():
    """Current job record, or None. Never raises."""
    try:
        with open(SYNC_JOB_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except (FileNotFoundError, ValueError, OSError):
        return None


def _sync_job_write(job):
    """Write the job record atomically so a concurrent poll never reads a
    half-written file."""
    try:
        os.makedirs(SYNC_JOB_DIR, exist_ok=True)
        tmp = SYNC_JOB_FILE + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(job, f)
        os.replace(tmp, SYNC_JOB_FILE)
    except OSError as e:
        print(f"[sync] could not write job file: {e}")


def _sync_job_is_stale(job):
    if not job or job.get("status") != "running":
        return False
    return (time.time() - float(job.get("heartbeat") or 0)) > SYNC_STALE_SECONDS


def _sync_percent(p):
    """Completion estimate, weighted by how long each phase actually takes.

    Bands: payments 1-5, transactions 5-17, registrations 18-40, photos 40-98.
    Registrations and photos both report incremental counts (see the
    [PROGRESS] lines in offline_sync.py), so within those two bands — which are
    almost all of the wall-clock time — the bar advances continuously.
    """
    phase = p.get("phase")

    def band(low, high, done, total):
        if not total:
            return low
        return low + int((high - low) * min(max(done / total, 0.0), 1.0))

    if phase in (None, "starting"):
        return 1
    if phase == "payments":
        return 4
    if phase == "transactions":
        return band(5, 17, p.get("paymentsDone") or 0, p.get("paymentsTotal") or 0)
    if phase == "registrations":
        return band(
            18, 40,
            p.get("registrationsDone") or 0,
            p.get("registrationsTotal") or p.get("beneficiaries") or 0,
        )
    if phase == "photos":
        return band(
            40, 98,
            p.get("photosDone") or 0,
            p.get("photosTotal") or p.get("beneficiaries") or 0,
        )
    if phase in ("done", "failed"):
        return 100
    return 4


def _sync_worker(job_id, program_id):
    """Run offline_sync.py, tailing its stdout to keep the job file current."""
    env = os.environ.copy()
    env["PROGRAM_ID"] = str(program_id)

    progress = {
        "phase": "starting",
        "paymentsTotal": None,
        "paymentsDone": 0,
        "transactions": 0,
        "beneficiaries": None,
        "registrationsFetched": 0,
        "registrationsFailed": 0,
        "registrationsDone": 0,
        "registrationsTotal": 0,
        "photosDone": 0,
        "photosTotal": 0,
    }
    job = {
        "jobId": job_id,
        "programId": str(program_id),
        "status": "running",
        "startedAt": time.time(),
        "heartbeat": time.time(),
        "progress": progress,
        "percent": 2,
        "message": "",
        "error": "",
    }
    _sync_job_write(job)

    stdout_tail = _sync_deque(maxlen=300)
    stderr_chunks = []
    final_line = ""

    def _flush(force=False):
        now = time.time()
        if force or now - _flush.last >= 1.0:
            _flush.last = now
            job["heartbeat"] = now
            job["percent"] = _sync_percent(progress)
            _sync_job_write(job)

    _flush.last = 0.0

    try:
        proc = subprocess.Popen(
            [sys.executable, "offline_sync.py"],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
            errors="replace",
            env=env,
            bufsize=1,
        )
    except Exception as e:
        job["status"] = "failed"
        job["error"] = f"Could not start offline_sync.py: {e}"
        job["finishedAt"] = time.time()
        job["heartbeat"] = time.time()
        job["percent"] = 100
        _sync_job_write(job)
        return

    # Drain stderr on its own thread. Reading it only after stdout closes would
    # deadlock if a traceback filled the 64KB pipe buffer.
    def _drain_stderr():
        try:
            for line in proc.stderr:
                stderr_chunks.append(line)
        except Exception:
            pass

    err_thread = threading.Thread(target=_drain_stderr, daemon=True)
    err_thread.start()

    try:
        for raw in proc.stdout:
            line = raw.rstrip()
            stdout_tail.append(line)

            # Incremental counters first — these are the highest-frequency
            # lines and drive most of the bar's movement.
            m = _RE_PROGRESS.search(line)
            if m:
                kind, done, total = m.group(1), int(m.group(2)), int(m.group(3))
                if kind == "registrations":
                    progress["phase"] = "registrations"
                    progress["registrationsDone"] = done
                    progress["registrationsTotal"] = total
                else:
                    progress["phase"] = "photos"
                    progress["photosDone"] = done
                    progress["photosTotal"] = total
                _flush()
                continue

            m = _RE_OPEN_PAYMENTS.search(line)
            if m:
                progress["phase"] = "transactions"
                progress["paymentsTotal"] = int(m.group(1))
                _flush(force=True)
                continue

            m = _RE_PAYMENT_DONE.search(line)
            if m:
                progress["phase"] = "transactions"
                progress["paymentsDone"] += 1
                progress["transactions"] += int(m.group(2))
                _flush()
                continue

            m = _RE_BENEFICIARIES.search(line)
            if m:
                progress["phase"] = "registrations"
                progress["beneficiaries"] = int(m.group(1))
                _flush(force=True)
                continue

            m = _RE_REGISTRATIONS.search(line)
            if m:
                progress["phase"] = "photos"
                progress["registrationsFetched"] = int(m.group(2))
                progress["registrationsFailed"] = int(m.group(3))
                _flush(force=True)
                continue

            if _RE_PHOTO_OK.search(line):
                # Fallback only: if this app is ever run against an older
                # offline_sync.py that has no [PROGRESS] lines, count the
                # per-photo "[OK]" lines instead. When [PROGRESS] is present it
                # is authoritative and this must not double-count.
                progress["phase"] = "photos"
                if not progress.get("photosTotal"):
                    progress["photosDone"] += 1
                _flush()
                continue

            if _RE_FINAL.search(line):
                final_line = line
                _flush(force=True)
                continue

            if line.startswith("[INFO] Program ") or line.startswith("[INFO] Loaded "):
                progress["phase"] = "payments"
                _flush()
    except Exception as e:
        print(f"[sync] stdout reader error: {e}")

    proc.wait()
    err_thread.join(timeout=5)

    stdout_text = "\n".join(stdout_tail)
    stderr_text = "".join(stderr_chunks)

    print("\n[DEBUG] SYNC STDOUT (tail):\n", stdout_text)
    print("\n[DEBUG] SYNC STDERR:\n", stderr_text)

    if proc.returncode != 0:
        progress["phase"] = "failed"
        job["status"] = "failed"
        # offline_sync.py raises IncompleteSyncError (strict mode) rather than
        # writing a batch it knows is missing people. Surface that verbatim —
        # it is the most important message this app can show an FSP.
        detail = (stderr_text or stdout_text or "").strip()
        job["error"] = detail[-2000:] if detail else f"Sync exited with code {proc.returncode}"
        job["message"] = "Sync failed — no batch was written"
    else:
        progress["phase"] = "done"
        job["status"] = "done"
        if not final_line:
            for line in reversed(stdout_tail):
                if "beneficiaries" in line.lower():
                    final_line = line
                    break
        # Plain text: the page renders state with an icon and colour, so a
        # status glyph in the string would be redundant (and looks amateurish).
        job["message"] = final_line.strip() if final_line else "Sync completed"
        job["count"] = int(_RE_FINAL.match(job["message"]).group(1)) if _RE_FINAL.match(job["message"]) else None

    job["finishedAt"] = time.time()
    job["heartbeat"] = time.time()
    job["percent"] = 100
    _sync_job_write(job)


@app.route("/sync-fsp")
def sync_fsp():
    """Start a sync. Returns immediately; poll /sync-status for progress.

    Single-flight: if a sync is already running, its job is returned instead of
    launching a second offline_sync.py (two would race on get_next_batch_dir and
    both write batch directories).
    """
    program_id = session.get("fsp_program_id")
    if not program_id:
        return jsonify({"success": False, "message": "❌ No program selected"})

    with _sync_job_lock:
        existing = _sync_job_read()
        if existing and existing.get("status") == "running" and not _sync_job_is_stale(existing):
            return jsonify(
                {
                    "success": True,
                    "started": False,
                    "alreadyRunning": True,
                    "jobId": existing.get("jobId"),
                    "job": existing,
                    "message": "Sync already running",
                }
            )

        job_id = _uuid_mod.uuid4().hex[:12]
        _sync_job_write(
            {
                "jobId": job_id,
                "programId": str(program_id),
                "status": "running",
                "startedAt": time.time(),
                "heartbeat": time.time(),
                "progress": {"phase": "starting"},
                "percent": 1,
                "message": "",
                "error": "",
            }
        )

    threading.Thread(
        target=_sync_worker, args=(job_id, program_id), daemon=True
    ).start()

    return jsonify(
        {
            "success": True,
            "started": True,
            "jobId": job_id,
            "message": "Sync started",
        }
    )


@app.route("/sync-status")
def sync_status():
    """Current sync job state for the page's progress panel."""
    job = _sync_job_read()
    if not job:
        return jsonify({"status": "idle"})

    if _sync_job_is_stale(job):
        job = dict(job)
        job["status"] = "failed"
        job["error"] = (
            "The sync stopped reporting progress (the server process was "
            "restarted). Check the offline-cache folder, then sync again."
        )
        job["message"] = "❌ Sync interrupted."
    return jsonify(job)


@app.route("/fsp-logout")
def fsp_logout():
    # On a shared tablet, leaving the token behind would attribute the next
    # person's distribution to the previous user - a confidently wrong audit
    # trail, which is worse than an obviously shared one.
    _clear_login_session("fsp")
    return redirect(url_for("fsp_login"))


@app.route("/scan")
def scan():
    # Only FSP-logged-in users should scan
    lang = request.args.get("lang", "en")

    if not session.get("fsp_logged_in"):
        return redirect(url_for("fsp_login", lang=lang))

    username = session.get("fsp_username", "User")
    program_id = session.get("fsp_program_id")

    # Per-program scan method: QR enabled by default unless explicitly disabled
    qr_enabled = True
    try:
        full_display = load_display_config() or {}
        programs_map = full_display.get("programs", {}) or {}
        program_config = programs_map.get(str(program_id), {}) if program_id else {}
        qr_cfg = program_config.get("qr", {})
        # Default ON: only False when explicitly set to disabled
        qr_enabled = qr_cfg.get("enabled", True) if isinstance(qr_cfg, dict) else True
    except Exception as e:
        print(f"[scan] Failed to load qr config: {e}")
        qr_enabled = True

    # Fetch 121 program title
    program_title = ""
    if program_id:
        system_config = load_config()
        url121 = system_config.get("url121")
        if url121:
            try:
                login_resp = requests.post(
                    f"{url121}/api/users/login",
                    json={
                        "username": system_config.get("username121", ""),
                        "password": system_config.get("password121", ""),
                    },
                    timeout=8,
                )
                if login_resp.status_code == 201:
                    token = login_resp.json().get("access_token_general")
                    r = requests.get(
                        f"{url121}/api/programs/{program_id}",
                        cookies={"access_token_general": token},
                        timeout=8,
                    )
                    if r.status_code == 200:
                        titles = r.json().get("titlePortal", {})
                        program_title = titles.get(lang) or next(
                            iter(titles.values()), ""
                        )
            except Exception as e:
                print(f"[scan] Failed to fetch 121 program title: {e}")

    # The scan box can now match on a configured registration attribute, whose
    # value is Fernet-encrypted in the offline cache — so this page needs the
    # key, exactly as /beneficiary-offline does.
    #
    # CAVEAT: /scan is precached cacheFirst by the service worker, so this
    # baked-in value freezes at precache time and would go stale if
    # ENCRYPTION_KEY were rotated. scan.html prefers meta.encryptionKey from
    # IndexedDB and treats this only as the fallback.
    try:
        fernet_key = load_config().get("ENCRYPTION_KEY", "")
    except Exception:
        fernet_key = ""

    return render_template(
        "scan.html",
        lang=lang,
        t=translations.get(lang, translations["en"]),
        username=username,
        program_title=program_title,
        program_id=program_id or "",
        qr_enabled=qr_enabled,
        fernet_key=fernet_key,
    )

@app.route("/service-worker.js")
def sw():
    return send_from_directory(
        "static", "service-worker.js", mimetype="application/javascript"
    )


@app.route("/manifest.webmanifest")
def manifest():
    return send_from_directory(
        "static", "manifest.webmanifest", mimetype="application/manifest+json"
    )


from io import BytesIO


@app.route("/api/offline/latest.zip")
def api_offline_latest_zip():
    base_dir = "offline-cache"
    if not os.path.isdir(base_dir):
        return jsonify({"error": "No offline cache found"}), 404

    # Filter by programId if provided
    program_id_filter = request.args.get("programId")

    batch_dirs = []
    for d in os.listdir(base_dir):
        full_path = os.path.join(base_dir, d)
        if not os.path.isdir(full_path):
            continue
        # Check batch_info.json for programId match
        if program_id_filter:
            batch_info_path = os.path.join(full_path, "batch_info.json")
            if os.path.exists(batch_info_path):
                try:
                    with open(batch_info_path) as f:
                        batch_info = json.load(f)
                    if str(batch_info.get("programId")) != str(program_id_filter):
                        continue
                except Exception:
                    pass
        batch_dirs.append(full_path)

    if not batch_dirs:
        return jsonify({"error": "No batches found for this program"}), 404

    latest = max(batch_dirs, key=os.path.getmtime)

    # Zip the latest batch in memory
    mem = BytesIO()
    with zipfile.ZipFile(mem, "w", zipfile.ZIP_DEFLATED) as zf:
        for root, _, files in os.walk(latest):
            for fname in files:
                full_path = os.path.join(root, fname)
                arcname = os.path.relpath(
                    full_path, latest
                )  # keep paths relative to batch root
                zf.write(full_path, arcname)
    mem.seek(0)

    return send_file(
        mem,
        mimetype="application/zip",
        as_attachment=True,
        download_name="latest_offline_cache.zip",
    )

@app.route("/api/program-field-uniqueness/<program_id>")
def api_program_field_uniqueness(program_id):
    """Is <field> unique across this programme's registrations in 121?

    Answers the config page's warning badge. Deliberately returns counts only —
    never the offending values — so this endpoint cannot be used to enumerate
    the caseload.

    Counts are a snapshot. A field that is unique today can gain a duplicate
    tomorrow, which is why offline_sync.py re-checks the same thing against the
    actual batch at sync time (Patch 4).
    """
    if not session.get("admin_logged_in"):
        return jsonify({"error": "unauthorized"}), 401

    field = (request.args.get("field") or "").strip()

    # The reference ID is unique by construction in 121 — no need to walk the
    # caseload to prove it.
    if not field or field == "referenceId":
        return jsonify({
            "field": field or "referenceId",
            "checked": 0, "blank": 0,
            "duplicateValues": 0, "affectedRegistrations": 0,
            "unique": True, "skipped": True,
        })

    system_config = load_config()
    url121 = system_config.get("url121")
    if not url121:
        return jsonify({"error": "121 URL is not configured"}), 503

    try:
        login_resp = requests.post(
            f"{url121}/api/users/login",
            json={
                "username": system_config.get("username121", ""),
                "password": system_config.get("password121", ""),
            },
            timeout=10,
        )
        if login_resp.status_code != 201:
            return jsonify({"error": "121 login failed"}), 502
        cookies = {
            "access_token_general": login_resp.json().get("access_token_general")
        }

        # Same pagination discipline as offline_sync._fetch_all: ask for
        # limit=-1 first, then verify against meta.totalItems and page
        # explicitly if the server paginated us anyway. A silently truncated
        # read here would report "unique" for a field that is not.
        url = f"{url121}/api/programs/{program_id}/registrations"
        rows, total = [], None

        r = requests.get(url, cookies=cookies, params={"limit": -1}, timeout=60)
        r.raise_for_status()
        payload = r.json()
        rows = payload.get("data", payload) if isinstance(payload, dict) else payload
        rows = rows if isinstance(rows, list) else []
        meta = payload.get("meta") if isinstance(payload, dict) else None
        total = (meta or {}).get("totalItems")

        if total is not None and len(rows) < total:
            rows = []
            page, page_size, max_pages = 1, 1000, 100
            while page <= max_pages:
                pr = requests.get(
                    url, cookies=cookies,
                    params={"limit": page_size, "page": page}, timeout=60,
                )
                pr.raise_for_status()
                pp = pr.json()
                batch = pp.get("data", []) if isinstance(pp, dict) else pp
                if not batch:
                    break
                rows.extend(batch)
                if total is not None and len(rows) >= total:
                    break
                page += 1

        # Truthful about incompleteness rather than quietly reporting on a
        # partial read.
        truncated = total is not None and len(rows) < total

        from collections import Counter
        values = [str((row or {}).get(field) or "").strip() for row in rows]
        blank = sum(1 for v in values if not v)
        counts = Counter(v.lower() for v in values if v)
        dupes = {v: c for v, c in counts.items() if c > 1}

        return jsonify({
            "field": field,
            "checked": len(rows),
            "reportedTotal": total,
            "truncated": truncated,
            "blank": blank,
            "duplicateValues": len(dupes),
            "affectedRegistrations": sum(dupes.values()),
            "unique": not dupes,
        })

    except Exception as e:
        print(f"[uniqueness] program {program_id} field '{field}': {e}")
        return jsonify({"error": str(e)}), 502


@app.route("/ping")
def ping():
    return "ok", 200


@app.route("/beneficiary-offline")
def beneficiary_offline():
    # expected: /beneficiary-offline?uuid=<registrationReferenceId>&lang=en&program_id=<id>
    uuid = request.args.get("uuid")
    lang = request.args.get("lang", session.get("lang", "en"))
    session["lang"] = lang
    # CHANGE: don't return 400; render a shell so the SW can precache a 200
    if not uuid:
        uuid = ""

    # Prefer URL param, fall back to session (handles SW precache and fresh tabs)
    program_id = request.args.get("program_id") or session.get("fsp_program_id", "")

    # load display config scoped to the active program
    try:
        full_config = load_display_config()
        programs_map = full_config.get("programs", {}) or {}
        # Empty dict fallback (NOT the root config) — the root has no 'photo' key
        # so falling back to it would silently disable the photo section.
        program_config = programs_map.get(str(program_id), {}) if program_id else {}
        display_fields = program_config.get("fields", [])
        photo_config = program_config.get("photo", {})
    except Exception:
        display_fields = []
        photo_config = {}

    config = load_config()
    enc_key = config.get("ENCRYPTION_KEY", "")
    # Legacy template variable. This page is served cacheFirst by the service
    # worker, so anything baked in here is frozen at precache time — it must
    # never be the authority. The device resolves the column per tranche.
    _cols = set(get_fsp_column_map(program_id).values())
    column_to_match = _cols.pop() if len(_cols) == 1 else ""
    return render_template(
        "beneficiary_offline.html",
        uuid=uuid,  # may be "", the page will prefer URL ?uuid=...
        lang=lang,
        t=translations.get(lang, translations["en"]),
        display_fields=display_fields,
        photo_config=photo_config,
        fernet_key=enc_key,
        column_to_match=column_to_match,
        program_id=program_id,
        program_currency=config.get("programCurrency", ""),
    )


@app.route("/success-offline")
def success_offline():
    lang = request.args.get("lang", "en")
    t = translations.get(lang, translations["en"])
    program_id = request.args.get("program_id") or session.get("fsp_program_id", "")
    return render_template(
        "success_offline.html", lang=lang, t=t, program_id=program_id
    )


@app.route("/system-config.json")
def system_config_json():
    config = load_config()
    column = config.get("COLUMN_TO_MATCH")
    if not column:
        return jsonify({"error": "COLUMN_TO_MATCH missing"}), 500

    return jsonify({"COLUMN_TO_MATCH": column})


@app.route("/api/match-columns/<program_id>")
def api_match_columns(program_id):
    """Diagnostic: the full FSP-configuration -> columnToMatch mapping, plus the
    resolved column for each open payment.

    This is the view that did not exist when a single wrong column silently
    reached the field. Open it before any distribution to confirm the payment
    being run resolves to the column you expect.
    """
    token = get_121_token()
    if not token:
        return jsonify({"error": "Login to 121 failed"}), 502

    col_map = get_fsp_column_map(program_id, token=token)
    if not col_map:
        return jsonify({"error": "No FSP configuration reports a columnToMatch"}), 404

    payments = []
    try:
        config = load_config()
        r = requests.get(
            f"{config['url121']}/api/programs/{program_id}/payments",
            cookies={"access_token_general": token},
            params={"limit": -1},
            timeout=15,
        )
        if r.status_code == 200:
            payload = r.json()
            items = payload if isinstance(payload, list) else payload.get("data", [])
            for p in items:
                pid = p.get("paymentId", p.get("id"))
                if pid is None:
                    continue
                payments.append({
                    "paymentId": pid,
                    "name": p.get("name"),
                    "waiting": ((p.get("aggregationsPerStatus") or {}).get("waiting") or {}).get("count", 0),
                    "columnToMatch": get_column_for_payment(
                        program_id, pid, col_map=col_map, token=token
                    ),
                })
    except Exception as e:
        print(f"[api_match_columns] {e}")

    return jsonify({
        "programId": str(program_id),
        "fspConfigurations": col_map,
        "payments": payments,
        "ambiguous": len(set(col_map.values())) > 1,
    })


def get_fsp_column_map(program_id, token=None):
    """{programFspConfigurationName: columnToMatch}, or {} on failure.

    Mirrors offline_sync.build_fsp_column_map. Keyed on the configuration NAME
    — never array order.
    """
    config = load_config()
    url121 = config.get("url121")
    if not (url121 and program_id):
        return {}

    token = token or get_121_token()
    if not token:
        return {}

    try:
        r = requests.get(
            f"{url121}/api/programs/{program_id}/fsp-configurations",
            cookies={"access_token_general": token},
            timeout=10,
        )
        if r.status_code != 200:
            print(f"[get_fsp_column_map] program {program_id}: HTTP {r.status_code}")
            return {}
        payload = r.json()
        configs = payload if isinstance(payload, list) else payload.get("data", [])
    except Exception as e:
        print(f"[get_fsp_column_map] API error: {e}")
        return {}

    col_map = {}
    for fsp in configs:
        if not isinstance(fsp, dict):
            continue
        name = fsp.get("name")
        if not name:
            continue
        for prop in fsp.get("properties", []):
            if prop.get("name") == "columnToMatch" and prop.get("value"):
                col_map[name] = prop["value"]
                break
    return col_map


def get_column_for_payment(program_id, payment_id, col_map=None, token=None):
    """The columnToMatch for ONE payment, or None.

    This is the ONLY correct scope. There is deliberately no program-level
    fallback and no global: a program can run several FSP configurations at
    once with different columns, and picking one arbitrarily silently
    reconciles against the wrong field.
    """
    config = load_config()
    url121 = config.get("url121")
    if not (url121 and program_id and payment_id):
        return None

    token = token or get_121_token()
    if not token:
        return None
    if col_map is None:
        col_map = get_fsp_column_map(program_id, token=token)
    if not col_map:
        return None

    try:
        r = requests.get(
            f"{url121}/api/programs/{program_id}/payments/{payment_id}",
            cookies={"access_token_general": token},
            timeout=10,
        )
        if r.status_code != 200:
            print(f"[get_column_for_payment] payment {payment_id}: HTTP {r.status_code}")
            return None
        fsps = r.json().get("fsps") or []
    except Exception as e:
        print(f"[get_column_for_payment] API error: {e}")
        return None

    names = [
        f.get("programFspConfigurationName")
        for f in fsps
        if isinstance(f, dict) and f.get("programFspConfigurationName")
    ]
    columns = {col_map[n] for n in names if n in col_map}

    if len(columns) == 1:
        return columns.pop()
    if len(columns) > 1:
        print(
            f"[get_column_for_payment] payment {payment_id} spans FSP configs "
            f"{names} with conflicting columns {sorted(columns)} — cannot reconcile."
        )
    return None


def get_121_token():
    import requests

    config = load_config()
    username = config.get("username121")
    password = config.get("password121")
    base_url = config.get("url121")

    if not username or not password or not base_url:
        print("❌ Missing 121 credentials")
        return None

    login_url = f"{base_url}/api/users/login"

    try:
        resp = requests.post(
            login_url, json={"username": username, "password": password}, timeout=8
        )

        if resp.status_code != 201:
            print(f"❌ Login failed ({resp.status_code}): {resp.text}")
            return None

        # 121 API returns token in JSON (correct behaviour)
        token = resp.json().get("access_token_general")
        if not token:
            print("❌ Login succeeded but no token returned")
            return None

        return token

    except Exception as e:
        print(f"❌ 121 API error: {e}")
        return None


# ---------------------------------------------------------------------------
# OPERATOR ATTRIBUTION
#
# Both login routes already authenticate the person against the 121 API
# (POST /api/users/login, accepted on 201) - there is no local password check.
# They used to throw away the access_token_general that 121 hands back, so every
# subsequent call went through get_121_token() as the shared service account
# (USERNAME_121 / PASSWORD_121 in Azure App Settings). In 121's audit log that
# made every payment reconciliation look like one robot user.
#
# Keeping the token means 121 records the real person against the payment.
# No password is ever stored - only the token 121 itself issued.
#
# Reads (programme titles, registration attributes, columnToMatch) and the
# background sync in offline_sync.py deliberately keep using the service
# account: there is no human behind them, and a field user may legitimately
# lack program.read, which would turn working dropdowns into empty ones.
# ---------------------------------------------------------------------------
def _store_login_session(role, username, response):
    """Record who logged in, plus the 121 token issued to them."""
    token = None
    permissions = None
    try:
        body = response.json()
        token = body.get("access_token_general")
        perms = body.get("permissions")
        if isinstance(perms, dict):
            permissions = {str(k): v for k, v in perms.items()}
    except Exception:
        pass

    # Some 121 builds set the token as a cookie rather than in the body.
    if not token:
        try:
            token = response.cookies.get("access_token_general")
        except Exception:
            token = None

    session[role + "_logged_in"] = True
    session[role + "_username"] = username
    session[role + "_token"] = token
    session[role + "_permissions"] = permissions
    session[role + "_login_at"] = time.time()

    if not token:
        print(
            "[auth] WARNING: 121 login for %s returned no access_token_general; "
            "this session's submissions will fall back to the service account."
            % username
        )


def _clear_login_session(role):
    """Wipe everything identifying this operator."""
    for suffix in ("_logged_in", "_username", "_token", "_permissions", "_login_at"):
        session.pop(role + suffix, None)
    if role == "fsp":
        session.pop("fsp_program_id", None)


# Permissions 121 requires for PATCH /api/programs/{id}/registrations, which is
# what /submit-registration-updates calls to push editable-field values.
#
# Both are required rather than either: 121 splits ordinary registration
# attributes (attribute.update) from PII-flagged ones (personal.update), the
# app cannot tell which category a configured editable field falls into without
# an extra per-attribute lookup, and the bulk PATCH is all-or-nothing — one
# missing permission fails the entire batch, not just the PII rows.
REGISTRATION_UPDATE_PERMISSIONS = (
    "registration:attribute.update",
    "registration:personal.update",
)


def _has_program_permission(role, program_id, permission):
    """Tri-state: True / False / None when 121 did not tell us.

    Only a definite False is acted on. If the login response had no readable
    permissions map, we return None and let the request proceed - 121 is the
    real gate, and blocking a whole distribution because a response shape
    changed would be far worse than a late 403.
    """
    perms = session.get(role + "_permissions")
    if not isinstance(perms, dict):
        return None
    entry = perms.get(str(program_id))
    if isinstance(entry, dict):
        entry = entry.get("permissions") or entry.get("permission")
    if not isinstance(entry, (list, tuple, set)):
        return None
    return permission in entry


def get_121_token_for_request(role="fsp"):
    """(token, actor_label, is_user_attributed) for a write made by a human.

    Falls back to the service account when there is no operator token - an old
    session, or a 121 build that stops returning the token - so a distribution
    degrades to the previous behaviour rather than stopping.
    """
    token = session.get(role + "_token")
    username = session.get(role + "_username")
    if token:
        return token, (username or role), True

    print(
        "[auth] NOTE: no operator token in session, so this submission is "
        "attributed to the service account."
    )
    return get_121_token(), "the shared 121 service account", False


@app.route("/submit-payments", methods=["POST"])
def submit_payments():
    import csv
    import io
    import os
    import json
    import traceback
    from datetime import datetime
    from cryptography.fernet import Fernet, InvalidToken

    try:
        # Load config
        config = load_config()
        # Use active program from session (multi-program support)
        program_id = session.get("fsp_program_id") or config.get("programId")
        fernet_key = config.get("ENCRYPTION_KEY")

        if not program_id:
            return (
                "❌ No active program selected. Please go back and select a program.",
                400,
            )

        # Match columns are per FSP configuration, reached via the payment.
        # Built once here; each paymentId is resolved against it in the
        # submission loop below. No token is passed: this is a configuration
        # read, so it uses the service account, which is guaranteed to have
        # access even when the submitting user's role does not.
        fsp_col_map = get_fsp_column_map(program_id)
        if not fsp_col_map:
            return (
                f"❌ Could not read the FSP configurations for program {program_id}. "
                "Check 121 connectivity and that this account has access.",
                400,
            )
        payment_columns = {}  # paymentId -> column, resolved lazily below

        if not fernet_key:
            return "❌ Missing ENCRYPTION_KEY in system_config.json", 400

        # Fernet decryptor
        try:
            fernet = Fernet(fernet_key.encode())
        except ValueError as e:
            return f"❌ Invalid Fernet key: {e}", 400

        # Get uploaded CSV file
        if "csv" not in request.files:
            return "❌ No CSV file provided", 400

        file = request.files["csv"]
        if file.filename == "":
            return "❌ Empty filename", 400

        try:
            csv_content = file.stream.read().decode("utf-8")
        except Exception as e:
            return f"❌ Failed to read CSV: {e}", 400

        reader = csv.DictReader(io.StringIO(csv_content))
        rows = list(reader)

        if not rows:
            return "❌ CSV is empty", 400

        # -------------------------------
        # LOAD OFFLINE CACHE FOR PAYMENT MAPPING
        # -------------------------------
        cache_base = "offline-cache"

        import re

        def extract_batch_number(name):
            match = re.search(r"payment-recent-batch-(\d+)", name)
            return int(match.group(1)) if match else -1

        # Defensive: cache_base may not exist on a fresh deployment
        if not os.path.isdir(cache_base):
            return (
                f"❌ Offline cache directory '{cache_base}' not found on the server. "
                "Run sync first to create it.",
                400,
            )

        # Filter batch dirs to only those belonging to the active program
        all_dirs = [
            d for d in os.listdir(cache_base) if d.startswith("payment-recent-batch-")
        ]
        program_dirs = []
        for d in all_dirs:
            batch_info_path = os.path.join(cache_base, d, "batch_info.json")
            if os.path.exists(batch_info_path):
                try:
                    with open(batch_info_path) as f:
                        bi = json.load(f)
                    if str(bi.get("programId")) == str(program_id):
                        program_dirs.append(d)
                except Exception:
                    pass
            else:
                program_dirs.append(d)  # include legacy batches without batch_info

        batch_dirs = sorted(program_dirs, key=extract_batch_number)

        if not batch_dirs:
            return (
                "❌ No recent payment batches found for this program — run sync first.",
                400,
            )

        latest_batch = batch_dirs[-1]
        print(f"[DEBUG] Using batch folder: {latest_batch}")

        latest_batch_info = {}
        latest_batch_info_path = os.path.join(
            cache_base, latest_batch, "batch_info.json"
        )
        if os.path.exists(latest_batch_info_path):
            try:
                with open(latest_batch_info_path, "r", encoding="utf-8") as f:
                    latest_batch_info = json.load(f)
            except Exception:
                latest_batch_info = {}

        # Legacy batches (without metadata) are assumed encrypted.
        cache_data_encrypted = bool(latest_batch_info.get("dataEncrypted", True))

        reg_cache_path = os.path.join(
            cache_base, latest_batch, "registrations_cache.json"
        )
        if not os.path.exists(reg_cache_path):
            return "❌ registrations_cache.json missing — run sync again.", 400

        try:
            with open(reg_cache_path, "r", encoding="utf-8") as f:
                reg_data = json.load(f)
        except Exception as e:
            return f"❌ Failed to load registrations_cache.json — {e}", 500

        # The old "match value -> paymentId" cache map lived here. It keyed rows
        # by match value under ONE program-wide column, which cannot work now
        # that different payments use different columns — it would reconcile an
        # arbitrary payment. paymentId is required on every row instead.
        #
        # -------------------------------
        # -------------------------------
        grouped = {}
        skipped_no_payment_id = 0

        # Current devices emit a column-neutral "matchValue". Older ones emit a
        # column named after whatever they thought the match column was, so
        # fall back to the first non-metadata column in the row.
        _META_COLS = {"status", "paymentId", "columnToMatch", "_scandroid_encrypted"}

        for row in rows:
            raw_value = (row.get("matchValue") or "").strip()
            if not raw_value:
                for _k, _v in row.items():
                    if _k and _k not in _META_COLS and _k != "matchValue" and _v:
                        raw_value = str(_v).strip()
                        break
            status = row.get("status", "").strip()

            # Only decrypt CSV values when explicitly marked as encrypted.
            row_value_encrypted = str(
                row.get("_scandroid_encrypted", "")
            ).strip().lower() in {
                "1",
                "true",
                "yes",
            }  # permissive “truthy” markers so different CSV exporters/users can flag encrypted rows without exact casing/format
            if row_value_encrypted:
                try:
                    raw_value = fernet.decrypt(raw_value.encode()).decode().strip()
                except InvalidToken as e:
                    print(f"[!] Failed to decrypt incoming match value — invalid token: {e}")
                    continue
            # Prefer the paymentId the FSP actually selected on the device.
            #
            # A beneficiary can hold several concurrent payments. match_to_pid is
            # keyed only by the match value, so it can hold ONE paymentId per
            # person — using it would reconcile an arbitrary payment (whichever
            # cache record was written last) instead of the one that was handed
            # out. The device now sends its choice in a paymentId column; the
            # cache lookup remains only as a fallback for CSVs produced by an
            # older client build.
            # paymentId is now REQUIRED. The old cache fallback keyed rows by
            # match value under a single program-wide column — which cannot work
            # once different payments use different columns, and would silently
            # close out the wrong payment. A row without one is rejected rather
            # than guessed at.
            payment_id = str(row.get("paymentId") or "").strip()
            if not payment_id:
                skipped_no_payment_id += 1
                continue

            grouped.setdefault(payment_id, []).append(
                {"matchValue": raw_value, "status": status}
            )

        if not grouped:
            if skipped_no_payment_id:
                return (
                    f"❌ No valid rows to submit — {skipped_no_payment_id} row(s) "
                    "carried no paymentId. This export came from an older app "
                    "version. Sync the device while online and send again.",
                    400,
                )
            return (
                "❌ No valid rows to submit — the CSV contained no usable match "
                "values. Re-sync and try again.",
                400,
            )

        if skipped_no_payment_id:
            print(f"[warn] {skipped_no_payment_id} row(s) skipped: no paymentId")
        # -------------------------------
        # SUBMIT TO 121 /paymentId/excel-reconciliation
        # -------------------------------
        token, actor_label, user_attributed = get_121_token_for_request("fsp")
        if not token:
            return "❌ Login to 121 failed", 401

        # Early warning only - see _has_program_permission. A definite "no" is
        # worth catching before we upload; anything less proceeds and lets 121
        # decide.
        if user_attributed and _has_program_permission(
            "fsp", program_id, "payment.update"
        ) is False:
            return (
                f"❌ Your 121 account ({actor_label}) is not allowed to reconcile "
                f"payments for program {program_id}. Ask your 121 administrator to "
                "add the 'payment.update' permission to your role.",
                403,
            )

        print(f"[auth] submitting payment reconciliation as {actor_label}")

        success_count = 0
        fail_count = 0
        failure_details = []

        for pid, items in grouped.items():
            # The header 121 expects is a property of THIS payment's FSP
            # configuration. Resolved per payment, live, inside the loop that
            # already groups by payment.
            if pid not in payment_columns:
                payment_columns[pid] = get_column_for_payment(
                    program_id, pid, col_map=fsp_col_map, token=token
                )
            pay_column = payment_columns[pid]

            if not pay_column:
                fail_count += 1
                failure_details.append(
                    f"paymentId {pid}: could not determine its match column from "
                    "the 121 FSP configuration — not submitted"
                )
                continue

            print(f"[OK] paymentId {pid}: reconciling on column '{pay_column}'")

            output_buffer = io.StringIO()
            writer = csv.DictWriter(
                output_buffer, fieldnames=[pay_column, "status"]
            )
            writer.writeheader()

            for item in items:
                writer.writerow(
                    {pay_column: item["matchValue"], "status": item["status"]}
                )

            upload_url = f"{config['url121']}/api/programs/{program_id}/payments/{pid}/excel-reconciliation"

            files = {
                "file": ("reconciliation.csv", output_buffer.getvalue(), "text/csv")
            }

            try:
                upload_resp = requests.post(
                    upload_url,
                    files=files,
                    cookies={"access_token_general": token},
                    timeout=30,
                )
            except requests.RequestException as e:
                fail_count += 1
                failure_details.append(f"paymentId {pid}: network error — {e}")
                print(f"[ERROR] Network error submitting paymentId {pid}: {e}")
                continue

            if upload_resp.status_code == 201:
                success_count += 1
                print(f"[OK] Submitted to paymentId {pid}")
            elif upload_resp.status_code == 401 and user_attributed:
                fail_count += 1
                failure_details.append(
                    f"paymentId {pid}: your 121 login has expired — log out, log "
                    "back in and send again. Nothing was lost."
                )
                print(f"[ERROR] Operator token rejected (401) for paymentId {pid}")
            else:
                fail_count += 1
                snippet = (upload_resp.text or "")[:300]
                failure_details.append(
                    f"paymentId {pid}: HTTP {upload_resp.status_code} — {snippet}"
                )
                print(
                    f"[ERROR] Failed to submit to paymentId {pid}: {upload_resp.status_code} — {upload_resp.text}"
                )

        # -------------------------------
        # FINAL RESPONSE
        # -------------------------------
        if success_count > 0:
            msg = f"✅ Submitted to {success_count} paymentId(s)."
            if fail_count:
                msg += f" ❌ {fail_count} failed: " + " | ".join(failure_details)
            return msg, 200
        else:
            detail = (
                " | ".join(failure_details)
                if failure_details
                else "no details available"
            )
            return f"❌ All submissions failed. {detail}", 502

    except Exception as e:
        # Last-resort handler: surface a readable message to the frontend
        # AND log the full traceback for server-side debugging.
        tb = traceback.format_exc()
        print(f"[FATAL] /submit-payments crashed: {e}\n{tb}")
        return (
            f"❌ Server error in /submit-payments: {type(e).__name__}: {e}",
            500,
        )


@app.route('/submit-registration-updates', methods=['POST'])
def submit_registration_updates():
    """Push editable-field updates (e.g. prepaid card numbers) captured during
    an offline distribution back to 121 via the bulk-update endpoint:

        PATCH /api/programs/{programId}/registrations   (multipart/form-data)

    The CSV must have a `referenceId` column plus one column per attribute to
    update. Empty cells are treated by 121 as "set to empty", so the frontend
    only emits cells that actually carry a value. This is separate from
    /submit-payments so the payment reconciliation flow stays untouched.
    """
    import io
    import csv as _csv
    import traceback

    try:
        config = load_config()
        program_id = session.get("fsp_program_id") or config.get("programId")
        if not program_id:
            return "❌ No active program selected. Please go back and select a program.", 400

        url121 = config.get("url121")
        if not url121:
            return "❌ Missing url121 in system_config.json", 400

        if 'csv' not in request.files:
            return "❌ No CSV file provided", 400

        file = request.files['csv']
        if file.filename == '':
            return "❌ Empty filename", 400

        try:
            csv_bytes = file.stream.read()
            csv_text = csv_bytes.decode("utf-8")
        except Exception as e:
            return f"❌ Failed to read CSV: {e}", 400

        # Light validation: must be non-empty and contain a referenceId column.
        reader = _csv.DictReader(io.StringIO(csv_text))
        headers = reader.fieldnames or []
        rows = list(reader)
        if not rows:
            return "❌ CSV is empty — nothing to update.", 400
        if "referenceId" not in headers:
            return "❌ CSV is missing the required 'referenceId' column.", 400
        if len(rows) > 100000:
            return "❌ Too many rows — 121 supports at most 100k rows per update.", 400

        token, actor_label, user_attributed = get_121_token_for_request("fsp")
        if not token:
            return "❌ Login to 121 failed", 401

        # Hiding the inputs in the UI is not enforcement: without this, an
        # account without the permission could POST here directly.
        #
        # payment.update is included because these edits are only ever sent as
        # part of a Send; the two registration permissions are what 121 itself
        # checks on the bulk PATCH. Tri-state — only a definite False is acted
        # on, so an unreadable permissions map still lets 121 be the judge.
        if user_attributed:
            missing = [
                perm
                for perm in ("payment.update",) + REGISTRATION_UPDATE_PERMISSIONS
                if _has_program_permission("fsp", program_id, perm) is False
            ]
            if missing:
                return (
                    f"❌ Your 121 account ({actor_label}) is not allowed to "
                    f"update registrations for program {program_id}. Ask your "
                    "121 administrator to add "
                    + ", ".join(f"'{p}'" for p in missing)
                    + " to your role.",
                    403,
                )

        print(f"[auth] submitting registration updates as {actor_label}")

        update_url = f"{url121}/api/programs/{program_id}/registrations"
        files = {"file": ("registration_updates.csv", csv_bytes, "text/csv")}
        # 121 requires a `reason` on every bulk update (stored in the audit log).
        # Naming the operator here puts the same attribution in the reason string
        # as on the token, which is what a reviewer actually reads in 121.
        data = {
            "reason": f"Updated during offline distribution (121 Scan) by {actor_label}"
        }

        try:
            resp = requests.patch(
                update_url,
                files=files,
                data=data,
                cookies={"access_token_general": token},
                timeout=60,
            )
        except requests.RequestException as e:
            return f"❌ Network error updating registrations: {e}", 502

        if resp.status_code in (200, 201, 202, 204):
            return f"✅ Updated {len(rows)} registration(s) in 121.", 200

        if resp.status_code == 401 and user_attributed:
            return (
                "❌ Your 121 login has expired. Please log out, log back in and "
                "try again.",
                401,
            )

        snippet = (resp.text or "")[:400]
        return (
            f"❌ 121 rejected the registration update (HTTP {resp.status_code}): {snippet}",
            502,
        )

    except Exception as e:
        tb = traceback.format_exc()
        print(f"[FATAL] /submit-registration-updates crashed: {e}\n{tb}")
        return f"❌ Server error in /submit-registration-updates: {type(e).__name__}: {e}", 500


@app.route("/invalid-qr")
def invalid_qr():
    # keep previously-selected language
    lang = request.args.get("lang", "en")
    reason = request.args.get("reason", "")

    program_id = request.args.get("program_id") or session.get("fsp_program_id", "")
    return render_template(
        "invalid-qr.html",
        reason=reason,
        lang=lang,
        t=translations.get(lang, translations["en"]),
        program_id=program_id,
    )


@app.route("/voucher-design", methods=["GET", "POST"])
def voucher_design():
    if not session.get("admin_logged_in"):
        return redirect(url_for("admin_login", lang=request.args.get("lang", "en")))

    lang = request.args.get("lang", "en")
    t = translations.get(lang, translations["en"])
    username = session.get("admin_username", "Admin")

    # ------------------------------------------------------
    # POST (SAVE A PROGRAM'S VOUCHER DESIGN) – multipart form
    # ------------------------------------------------------
    if request.method == "POST":
        program_id = request.form.get("programId")
        if not program_id:
            return jsonify({"success": False, "error": "Missing programId"}), 400

        existing = get_voucher_design(program_id) or {}
        design = {
            "title": (request.form.get("title") or "").strip(),
            "subtitle": (request.form.get("subtitle") or "").strip(),
            "logo1": existing.get("logo1"),
            "logo2": existing.get("logo2"),
            "logo1_size": _clean_size(request.form.get("logo1_size") or existing.get("logo1_size")),
            "logo2_size": _clean_size(request.form.get("logo2_size") or existing.get("logo2_size")),
            "show_qr": (request.form.get("show_qr", "1") != "0"),
            # The voucher-field order is owned by the Generate Vouchers page,
            # not this form — carry it through untouched or saving a design
            # would silently reset it.
            "fields": _voucher_field_config(existing),
        }

        # Explicit clears
        if request.form.get("logo1_clear") == "1":
            design["logo1"] = None
        if request.form.get("logo2_clear") == "1":
            design["logo2"] = None

        # New uploads (overwrite)
        try:
            f1 = request.files.get("logo1")
            if f1 and f1.filename:
                design["logo1"] = _save_logo_file(f1, program_id, "left")
            f2 = request.files.get("logo2")
            if f2 and f2.filename:
                design["logo2"] = _save_logo_file(f2, program_id, "right")
        except ValueError as e:
            return jsonify({"success": False, "error": str(e)}), 400

        try:
            save_voucher_design(program_id, design)
        except Exception as e:
            return jsonify({"success": False, "error": str(e)}), 500

        return jsonify({"success": True, "design": _design_for_client(design)})

    # ------------------------------------------------------
    # GET (LOAD PAGE)
    # ------------------------------------------------------
    try:
        system_config = load_config()
    except Exception:
        system_config = {}

    programs, _ = resolve_programs(system_config)
    designs = load_voucher_designs()

    active_program_id = request.args.get("programId")
    if not active_program_id and programs:
        active_program_id = programs[0]["id"]
    active_program_id = str(active_program_id) if active_program_id else None

    designs_client = {str(pid): _design_for_client(d) for pid, d in designs.items()}

    program_title = system_config.get("programTitle", "")

    return render_template(
        "voucher_design.html",
        programs=programs,
        active_program_id=active_program_id,
        designs=designs_client,
        program_title=program_title,
        lang=lang,
        username=username,
        t=t,
    )


@app.route("/vouchers", methods=["GET"])
def vouchers_page():
    if not session.get("admin_logged_in"):
        lang = request.args.get("lang", session.get("lang", "en"))
        return redirect(url_for("admin_login", lang=lang))

    # language
    lang = request.args.get("lang", session.get("lang", "en"))
    session["lang"] = lang
    t = translations.get(lang, translations["en"])

    # Load system config (for program title, currency, etc.)
    try:
        system_config = load_config()
    except Exception:
        system_config = {}

    # Username from session
    username = session.get("admin_username", "")

    # Program title from session OR fallback to system_config
    program_title = session.get("program_title") or system_config.get(
        "programTitle", ""
    )

    # Program dropdown + which programs already have a saved voucher design
    programs, _ = resolve_programs(system_config)
    designs = load_voucher_designs()
    designs_status = {str(p["id"]): (str(p["id"]) in designs) for p in programs}

    active_program_id = request.args.get("programId")
    if not active_program_id and programs:
        active_program_id = programs[0]["id"]
    active_program_id = str(active_program_id) if active_program_id else None

    return render_template(
        "vouchers.html",
        lang=lang,
        t=t,
        program_title=program_title,
        username=username,
        programs=programs,
        designs_status=designs_status,
        active_program_id=active_program_id,
    )


# ---------------------------------------------------------------------------
# Voucher source-file parsing
#
# The upload and download routes both need the same rows, and the download also
# needs the column ORDER as it appeared in the file (dict key order is not a
# reliable substitute once a saved field order is merged in). One parser keeps
# the two in step — they previously drifted, and the .xls branch of the
# download referenced an undefined `upload_path`, raising NameError on any .xls
# download.
#
# Returns (rows, columns):
#   rows    - [{lowercased_key: value}] with a normalised "referenceid" key
#   columns - [(lowercased_key, original_header_text)] in file order
# ---------------------------------------------------------------------------
REFERENCE_ID_ALIASES = (
    "referenceid",
    "reference id",
    "reference_id",
    "refid",
    "id",
)


def _normalise_header(raw):
    """(key, label) for one column header. Key is lowercased for lookups; the
    label keeps the author's original casing/script for printing."""
    label = str(raw).replace("﻿", "").strip()
    return label.lower(), label


def _resolve_reference_id(clean_row):
    for alias in REFERENCE_ID_ALIASES:
        value = clean_row.get(alias)
        if value not in (None, ""):
            return str(value).strip()
    return ""


def _parse_voucher_file(path):
    """Parse a voucher CSV/XLSX/XLS into (rows, columns).

    Raises ValueError for an unsupported extension.
    """
    name = str(path).lower()
    rows, columns = [], []

    # ---------------------------------------------------------------- CSV ----
    if name.endswith(".csv"):
        with open(path, "r", encoding="utf-8", errors="replace") as infile:
            reader = csv.DictReader(infile)
            for raw_header in reader.fieldnames or []:
                if raw_header is None:
                    continue
                key, label = _normalise_header(raw_header)
                if key and key not in dict(columns):
                    columns.append((key, label))

            for row in reader:
                clean_row = {}
                for raw_key, value in row.items():
                    if raw_key is None:
                        continue
                    key, _ = _normalise_header(raw_key)
                    clean_row[key] = value.strip() if isinstance(value, str) else value
                clean_row["referenceid"] = _resolve_reference_id(clean_row)
                if any(v not in (None, "") for v in clean_row.values()):
                    rows.append(clean_row)

    # --------------------------------------------------------------- XLSX ----
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)

        try:
            header_row = next(rows_iter)
        except StopIteration:
            header_row = []

        headers = []
        for raw_header in header_row:
            if raw_header in (None, ""):
                headers.append(None)
                continue
            key, label = _normalise_header(raw_header)
            headers.append(key or None)
            if key and key not in dict(columns):
                columns.append((key, label))

        for row in rows_iter:
            if not any(v not in (None, "") for v in row):
                continue
            clean_row = {}
            for key, value in zip(headers, row):
                if not key:
                    continue
                clean_row[key] = value.strip() if isinstance(value, str) else value
            clean_row["referenceid"] = _resolve_reference_id(clean_row)
            if any(v not in (None, "") for v in clean_row.values()):
                rows.append(clean_row)

    # ---------------------------------------------------------------- XLS ----
    elif name.endswith(".xls"):
        import xlrd

        wb = xlrd.open_workbook(path)
        sh = wb.sheet_by_index(0)

        # The first non-empty row is the real header row (some exports carry a
        # blank or title row above it).
        header_index = None
        for i in range(sh.nrows):
            if any(str(cell).strip() for cell in sh.row_values(i)):
                header_index = i
                break
        if header_index is None:
            return [], []

        headers = []
        for raw_header in sh.row_values(header_index):
            if raw_header in (None, ""):
                headers.append(None)
                continue
            key, label = _normalise_header(raw_header)
            headers.append(key or None)
            if key and key not in dict(columns):
                columns.append((key, label))

        for rx in range(header_index + 1, sh.nrows):
            values = sh.row_values(rx)
            if not any(str(v).strip() for v in values):
                continue
            clean_row = {}
            for key, value in zip(headers, values):
                if not key:
                    continue
                clean_row[key] = value.strip() if isinstance(value, str) else value
            clean_row["referenceid"] = _resolve_reference_id(clean_row)
            if any(v not in (None, "") for v in clean_row.values()):
                rows.append(clean_row)

    else:
        raise ValueError("Unsupported file type")

    return rows, columns


def _field_list_for_client(fields, rows):
    """Shape the merged field list for the reorder UI, adding a sample value
    taken from the first row that actually has one."""
    out = []
    for entry in fields:
        key = entry.get("key")
        sample = ""
        for row in rows:
            value = row.get(key)
            if value not in (None, ""):
                sample = str(value)
                break
        default_label = entry.get("default_label") or _default_label(key)
        out.append(
            {
                "key": key,
                "label": entry.get("label") or default_label,
                # The UI shows a reset affordance only for a renamed field, and
                # needs the default to reset back to.
                "default_label": default_label,
                "custom": bool(entry.get("custom")),
                "show": entry.get("show") is not False,
                "sample": sample[:60],
            }
        )
    return out


@app.route("/vouchers/upload", methods=["POST"])
def vouchers_upload():
    if not session.get("admin_logged_in"):
        return jsonify({"success": False, "message": "Not authorized"}), 403

    if "csv" not in request.files:
        return jsonify({"success": False, "message": "No file uploaded"}), 400

    f = request.files["csv"]
    filename = (f.filename or "").lower()

    try:
        os.makedirs("uploads", exist_ok=True)
        upload_path = os.path.join("uploads", filename)
        f.save(upload_path)

        try:
            rows, columns = _parse_voucher_file(upload_path)
        except ValueError:
            return jsonify({"success": False, "message": "Unsupported file type"}), 400

        program_id = request.form.get("program_id") or None

        # Reconcile the file's columns with this program's saved order, then
        # persist so the same order comes back on the next upload.
        design = get_voucher_design(program_id) if program_id else None
        merged = merge_voucher_fields(design, columns)

        if program_id:
            updated = dict(design or {})
            updated["fields"] = merged
            try:
                save_voucher_design(program_id, updated)
            except Exception as e:
                app.logger.warning("Could not save voucher field order: %s", e)

        # Save metadata
        session["voucher_file_path"] = upload_path
        session["voucher_count"] = len(rows)
        session["voucher_program_id"] = program_id

        return jsonify(
            {
                "success": True,
                "count": len(rows),
                "fields": _field_list_for_client(merged, rows),
            }
        )

    except Exception as e:
        return jsonify({"success": False, "message": f"Failed to parse file: {e}"}), 400


@app.route("/vouchers/fields", methods=["POST"])
def vouchers_fields():
    """Save the voucher field order / visibility for a program.

    Body: {"programId": "<id>", "fields": [{"key": ..., "label": ..., "show": bool}]}
    """
    if not session.get("admin_logged_in"):
        return jsonify({"success": False, "error": "Not authorized"}), 403

    payload = request.get_json(silent=True) or {}
    program_id = payload.get("programId") or session.get("voucher_program_id")
    if not program_id:
        return jsonify({"success": False, "error": "Missing programId"}), 400

    incoming = payload.get("fields")
    if not isinstance(incoming, list):
        return jsonify({"success": False, "error": "fields must be a list"}), 400

    fields, seen = [], set()
    for entry in incoming:
        if not isinstance(entry, dict):
            continue
        key = str(entry.get("key") or "").strip().lower()
        if not key or key in seen or _is_reference_id(key):
            continue
        seen.add(key)

        default_label = _clean_field_label(entry.get("default_label"), key)
        label = _clean_field_label(entry.get("label"), key, default_label)

        # "custom" is what protects a label from being refreshed out of the CSV
        # on the next upload. Trust it only when the label really does differ
        # from the default, so clearing the box reverts to the column name
        # rather than pinning the default forever.
        custom = bool(entry.get("custom")) and label != default_label

        stored = {"key": key, "label": label, "show": entry.get("show") is not False}
        if custom:
            stored["custom"] = True
        # default_label is derived on every upload — never persisted.
        fields.append(stored)

    design = dict(get_voucher_design(program_id) or {})
    design["fields"] = fields
    try:
        save_voucher_design(program_id, design)
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

    return jsonify({"success": True, "fields": fields})


@app.route("/vouchers/download", methods=["GET"])
def vouchers_download():
    if not session.get("admin_logged_in"):
        return redirect(url_for("admin_login", lang=request.args.get("lang", "en")))

    file_path = session.get("voucher_file_path")
    if not file_path or not os.path.exists(file_path):
        flash("No uploaded data to generate vouchers.", "error")
        return redirect(url_for("vouchers_page"))

    try:
        rows, _columns = _parse_voucher_file(file_path)
    except ValueError:
        flash("Unsupported voucher file type.", "error")
        return redirect(url_for("vouchers_page"))

    # -----------------------------
    # Generate PDF (using the selected program's design, if any)
    # -----------------------------
    program_id = session.get("voucher_program_id")
    design = get_voucher_design(program_id) if program_id else None

    pdf_io = generate_vouchers_pdf(
        rows,
        static_folder=os.path.join(app.root_path, "static"),
        design=design,
    )

    return send_file(
        pdf_io,
        mimetype="application/pdf",
        as_attachment=True,
        download_name="vouchers.pdf",
    )